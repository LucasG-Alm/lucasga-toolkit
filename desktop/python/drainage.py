"""Estudo de terreno por estruturas explícitas; sem dimensionamento hidráulico."""
from __future__ import annotations

import copy
import csv
import json
import math
import shutil
from pathlib import Path
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile

from pyproj import CRS, Transformer
from shapely.geometry import LineString, Point
from shapely.ops import substring

KML = 'http://www.opengis.net/kml/2.2'
NS = {'k': KML}
ET.register_namespace('', KML)


def metric_crs(epsg):
    crs = CRS.from_user_input(epsg)
    if not crs.is_projected or not crs.axis_info or any(abs(axis.unit_conversion_factor - 1) > 1e-9 for axis in crs.axis_info[:2]):
        raise ValueError('Use um CRS projetado em metros para medir o caminhamento.')
    return crs


def kml_document(source):
    source = Path(source)
    if source.suffix.lower() == '.kmz':
        with zipfile.ZipFile(source) as archive:
            names = [name for name in archive.namelist() if name.lower().endswith('.kml')]
            if 'doc.kml' in names:
                name = 'doc.kml'
            elif len(names) == 1:
                name = names[0]
            else:
                raise ValueError('KMZ precisa ter doc.kml ou um único KML principal.')
            if archive.getinfo(name).file_size > 50_000_000:
                raise ValueError('KML principal excede 50 MB.')
            return ET.fromstring(archive.read(name)), name
    if source.stat().st_size > 50_000_000:
        raise ValueError('KML excede 50 MB.')
    return ET.parse(source).getroot(), 'doc.kml'


def features(source, epsg, runtime=None, root=None):
    epsg = int(epsg)
    metric_crs(epsg)
    source = Path(source)
    paths, structures = [], []
    if source.suffix.lower() in {'.kml', '.kmz'}:
        tree, _ = kml_document(source)
        transform = Transformer.from_crs(4326, epsg, always_xy=True)
        for index, placemark in enumerate(tree.findall('.//k:Placemark', NS)):
            name = placemark.findtext('k:name', default='', namespaces=NS).strip()
            metadata = {item.get('name', '').lower(): item.findtext('k:value', default='', namespaces=NS)
                        for item in placemark.findall('k:ExtendedData/k:Data', NS)}
            for item in placemark.findall('k:ExtendedData/k:SchemaData/k:SimpleData', NS):
                metadata[item.get('name', '').lower()] = item.text or ''
            for line_index, line in enumerate(placemark.findall('.//k:LineString', NS)):
                coords = line.findtext('k:coordinates', default='', namespaces=NS)
                points = [tuple(map(float, coord.split(',')[:2])) for coord in coords.split()]
                if len(points) >= 2:
                    paths.append({'id': f'path:{index}:{line_index}', 'label': name or f'Caminhamento {index + 1}',
                                  'line': [transform.transform(*xy, errcheck=True) for xy in points]})
            for point_index, point in enumerate(placemark.findall('.//k:Point', NS)):
                raw = point.findtext('k:coordinates', default='', namespaces=NS).strip().split(',')
                if len(raw) >= 2 and name:
                    structures.append({'key': f'point:{index}:{point_index}', 'id': name,
                                       'xy': transform.transform(float(raw[0]), float(raw[1]), errcheck=True),
                                       'tipo': metadata.get('tipo', ''), 'diametro_mm': metadata.get('diametro_mm', metadata.get('dn', '')),
                                       'cota_terreno': metadata.get('cota_terreno', '')})
    else:
        import ezdxf
        import cad_geo
        with tempfile.TemporaryDirectory(prefix='toolkit-drenagem-leitura-') as work:
            doc = ezdxf.readfile(cad_geo.as_dxf(source, runtime or {}, root, work))
            for entity in doc.modelspace():
                kind, handle = entity.dxftype(), entity.dxf.handle
                if kind == 'LWPOLYLINE' and not entity.closed:
                    if any(abs(p[4]) > 1e-12 for p in entity.get_points()):
                        # A curved path must not silently become a shorter straight chord.
                        continue
                    if tuple(entity.dxf.extrusion) != (0, 0, 1):
                        continue
                    coords = [(p[0], p[1]) for p in entity.get_points()]
                    if len(coords) >= 2:
                        paths.append({'id': handle, 'label': f'{entity.dxf.layer} · {handle}', 'line': coords})
                elif kind == 'POLYLINE' and not entity.is_closed and (entity.is_2d_polyline or entity.is_3d_polyline):
                    if any(abs(vertex.dxf.get('bulge', 0)) > 1e-12 for vertex in entity.vertices):
                        continue
                    coords = [(v.dxf.location.x, v.dxf.location.y) for v in entity.vertices]
                    if len(coords) >= 2:
                        paths.append({'id': handle, 'label': f'{entity.dxf.layer} · {handle}', 'line': coords})
                elif kind in {'INSERT', 'POINT'}:
                    attrs = {attr.dxf.tag.lower(): attr.dxf.text for attr in entity.attribs} if kind == 'INSERT' else {}
                    point = entity.dxf.insert if kind == 'INSERT' else entity.dxf.location
                    identifier = attrs.get('id') or attrs.get('nome') or attrs.get('ponto')
                    if identifier:
                        structures.append({'key': handle, 'id': identifier, 'xy': (point.x, point.y),
                                           'tipo': attrs.get('tipo', ''), 'diametro_mm': attrs.get('diametro_mm', attrs.get('dn', '')),
                                           'cota_terreno': attrs.get('cota_terreno', '')})
    return {'paths': paths, 'structures': structures, 'epsg': epsg}


def inspect(source, epsg, runtime=None, root=None):
    data = features(source, epsg, runtime, root)
    return {'paths': [{'id': item['id'], 'label': item['label']} for item in data['paths']],
            'structures': [{'id': item['key'], 'label': item['id']} for item in data['structures']],
            'requests': math.ceil(len(data['structures']) / 100),
            'note': 'CAD: polilinhas abertas sem arcos; estruturas em blocos com ID/NOME/PONTO. KML: caminhos e pontos nomeados. Escolha explicitamente o eixo e as estruturas.'}


def read_selected(source, options, runtime=None, root=None):
    data = features(source, options.get('epsg') or 31983, runtime, root)
    selected = json.loads(options.get('structureIds') or '[]')
    path = next((item for item in data['paths'] if item['id'] == options.get('pathId')), None)
    if path is None or not isinstance(selected, list) or len(selected) < 2:
        raise ValueError('Selecione um caminhamento e pelo menos duas estruturas identificadas.')
    structures = [item for item in data['structures'] if item['key'] in selected]
    if len(structures) != len(set(selected)):
        raise ValueError('Uma estrutura selecionada não está mais na entrada. Inspecione novamente.')
    line = list(path['line'])
    if options.get('inverter'):
        line.reverse()
    return {'line': line, 'structures': structures, 'epsg': data['epsg'], 'source': str(source)}


def calculate(data, elevations, tolerance=2):
    line = LineString(data['line'])
    if not line.is_simple or line.length <= 0:
        raise ValueError('Caminhamento precisa ter extensão positiva e não cruzar a si mesmo.')
    if not math.isfinite(tolerance) or tolerance < 0:
        raise ValueError('Tolerância de afastamento inválida.')
    structures = []
    identifiers = set()
    if len(data['structures']) < 2 or len(elevations) != len(data['structures']):
        raise ValueError('São necessárias duas estruturas com cotas completas.')
    for structure, z in zip(data['structures'], elevations, strict=True):
        if structure['id'] in identifiers:
            raise ValueError(f"ID duplicado: {structure['id']}")
        identifiers.add(structure['id'])
        point = Point(structure['xy'])
        distance = line.distance(point)
        if distance > tolerance:
            raise ValueError(f"{structure['id']}: distância ao eixo {distance:.3f} m excede {tolerance} m.")
        if not math.isfinite(float(z)):
            raise ValueError('Cota não finita.')
        structures.append({**structure, 'estaca_m': line.project(point), 'cota_terreno_m': float(z), 'afastamento_m': distance})
    structures.sort(key=lambda item: item['estaca_m'])
    segments = []
    for a, b in zip(structures, structures[1:]):
        length = b['estaca_m'] - a['estaca_m']
        if length <= 1e-6:
            raise ValueError(f"{a['id']} e {b['id']} ocupam a mesma estaca.")
        delta = a['cota_terreno_m'] - b['cota_terreno_m']
        segments.append({'de': a['id'], 'ate': b['id'], 'extensao_m': length,
                         'cota_montante_m': a['cota_terreno_m'], 'cota_jusante_m': b['cota_terreno_m'],
                         'desnivel_m': delta, 'declividade_terreno_pct': delta / length * 100,
                         'diametro_mm': a.get('diametro_mm', '')})
    return {'structures': structures, 'segments': segments, 'fonte_cota': data.get('fonte_cota', 'Google Elevation — terreno estimado')}


def export_tables(result, folder):
    from openpyxl import Workbook
    folder = Path(folder)
    outputs = []
    groups = [('estruturas', [{
        'id': item['id'], 'tipo': item.get('tipo', ''), 'estaca_m': item['estaca_m'],
        'cota_terreno_m': item['cota_terreno_m'], 'diametro_mm': item.get('diametro_mm', ''),
        'x': item['xy'][0], 'y': item['xy'][1], 'afastamento_m': item['afastamento_m'], 'fonte_cota': result['fonte_cota'],
    } for item in result['structures']]), ('trechos', result['segments'])]
    book = Workbook()
    book.remove(book.active)
    for name, rows in groups:
        destination = folder / f'{name}.csv'
        with destination.open('x', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter=';')
            writer.writeheader()
            writer.writerows({key: "'" + value if isinstance(value, str) and value.startswith(('=', '+', '@')) else value for key, value in row.items()} for row in rows)
        outputs.append(destination)
        sheet = book.create_sheet(name.capitalize())
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append(list(row.values()))
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, (float, int)):
                    cell.number_format = '0.00'
        if name == 'trechos':
            for index in range(2, len(rows) + 2):
                sheet[f'F{index}'] = f'=D{index}-E{index}'
                sheet[f'G{index}'] = f'=F{index}/C{index}*100'
                sheet[f'F{index}'].number_format = sheet[f'G{index}'].number_format = '0.00'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    info = book.create_sheet('Leia-me')
    info.append(['Fonte das cotas', result['fonte_cota']])
    info.append(['Declividade', 'Do terreno, positiva quando desce no sentido do caminhamento. Não é declividade da tubulação.'])
    info.append(['Limite', 'Não dimensiona rede. Resolução da grade não garante precisão altimétrica. Tipos e diâmetros são os fornecidos.'])
    info.append(['Extensão analisada (m)', sum(item['extensao_m'] for item in result['segments'])])
    path = folder / 'estudo.xlsx'
    book.save(path)
    return [*outputs, path]


def write_kmz(data, result, target):
    source = Path(data['source']) if data.get('source') else None
    tree, main = kml_document(source) if source else (ET.Element(f'{{{KML}}}kml'), 'doc.kml')
    tree = copy.deepcopy(tree)
    document = tree.find('k:Document', NS)
    if document is None:
        document = ET.SubElement(tree, f'{{{KML}}}Document')
    def child(parent, name, text=None):
        element = ET.SubElement(parent, f'{{{KML}}}{name}')
        if text is not None:
            element.text = str(text)
        return element
    folder = child(document, 'Folder')
    child(folder, 'name', 'Estudo de terreno — resultados')
    transform = Transformer.from_crs(data['epsg'], 4326, always_xy=True)
    line = LineString(data['line'])
    for item in result['structures']:
        placemark = child(folder, 'Placemark')
        child(placemark, 'name', f"{item['id']} · CT {item['cota_terreno_m']:.2f} m")
        child(placemark, 'description', f"{result['fonte_cota']}; estaca {item['estaca_m']:.2f} m; tipo {item.get('tipo', '')}; diâmetro informado {item.get('diametro_mm', '')} mm")
        point = child(placemark, 'Point')
        child(point, 'altitudeMode', 'clampToGround')
        x, y = transform.transform(*item['xy'])
        child(point, 'coordinates', f'{x:.10f},{y:.10f},0')
    for index, segment in enumerate(result['segments']):
        a, b = result['structures'][index:index + 2]
        coords = list(substring(line, a['estaca_m'], b['estaca_m']).coords)
        placemark = child(folder, 'Placemark')
        child(placemark, 'name', f"{segment['de']} → {segment['ate']} · {segment['extensao_m']:.2f} m")
        child(placemark, 'description', f"declividade do terreno {segment['declividade_terreno_pct']:.4f}%; diâmetro informado {segment['diametro_mm']} mm; {result['fonte_cota']}")
        geometry = child(placemark, 'LineString')
        child(geometry, 'tessellate', 1)
        child(geometry, 'coordinates', ' '.join(f'{x:.10f},{y:.10f},0' for x, y in (transform.transform(*xy) for xy in coords)))
    with zipfile.ZipFile(target, 'x', zipfile.ZIP_DEFLATED) as archive:
        if source and source.suffix.lower() == '.kmz':
            with zipfile.ZipFile(source) as original:
                for info in original.infolist():
                    if info.filename != main:
                        archive.writestr(info, original.read(info))
        archive.writestr(main, ET.tostring(tree, encoding='utf-8', xml_declaration=True))


def write_dwg(data, result, target, runtime, root):
    import ezdxf
    import cad_geo
    sys.path.insert(0, str(root / 'ferramentas/proesplan'))
    import layouts
    sys.path.insert(0, str(root / 'ferramentas/galeria'))
    import galeria
    layer = 'DRENAGEM-TERRENO'
    lines = [f'(if (not (tblsearch "LAYER" "{layer}")) (command "_.-LAYER" "_N" "{layer}" "_C" "3" "{layer}" ""))']
    def text(x, y, value):
        lines.append(f'(entmake (list (cons 0 "TEXT") (cons 8 "{layer}") (cons 10 (list {x:.12g} {y:.12g} 0.0)) (cons 40 2.0) (cons 1 {galeria.lisp(value)})))')
    for item in result['structures']:
        text(*item['xy'], f"{item['id']} CT={item['cota_terreno_m']:.2f}m")
    line = LineString(data['line'])
    for index, segment in enumerate(result['segments']):
        a, b = result['structures'][index:index + 2]
        point = line.interpolate((a['estaca_m'] + b['estaca_m']) / 2)
        text(point.x, point.y + 3, f"{segment['de']}>{segment['ate']} L={segment['extensao_m']:.2f}m i-terreno={segment['declividade_terreno_pct']:.3f}%")
    with tempfile.TemporaryDirectory(prefix='toolkit-drenagem-cad-') as work:
        work = Path(work)
        before = ezdxf.readfile(cad_geo.as_dxf(data['source'], runtime, root, work))
        if any(e.dxf.layer == layer for e in before.modelspace()):
            raise ValueError('O desenho já tem DRENAGEM-TERRENO. Use a entrada original.')
        temporary_output = work / 'saida.dwg'
        layouts.roda_core(Path(runtime['core']), Path(data['source']), '\n'.join(lines) + '\n', temporary_output, work)
        after = ezdxf.readfile(layouts.dwg_para_dxf(Path(runtime['oda']), temporary_output, work))
        added = 0
        for name in before.layouts.names():
            old = layouts._mapa(before.layouts.get(name))
            new = layouts._mapa(after.layouts.get(name))
            if any(handle not in new or not layouts.igual(entity, new[handle]) for handle, entity in old.items()):
                raise RuntimeError('O DWG de saída alterou entidades existentes; não adote o resultado.')
            new_entities = [entity for handle, entity in new.items() if handle not in old]
            if any(entity.dxf.layer != layer for entity in new_entities):
                raise RuntimeError('O DWG de saída tem entidades novas fora da layer de resultados.')
            added += len(new_entities)
        if added != len(result['structures']) + len(result['segments']):
            raise RuntimeError('Quantidade de anotações do DWG não confere.')
        with temporary_output.open('rb') as origin, Path(target).open('xb') as destination:
            shutil.copyfileobj(origin, destination)


def execute(source, output, options, runtime, root, client):
    output = Path(output)
    if output.exists():
        raise FileExistsError('A pasta de saída já existe.')
    if Path(source).suffix.lower() in {'.dwg', '.dxf'} and options.get('crsConfirmed') is not True:
        raise ValueError('Confirme o CRS do desenho.')
    data = read_selected(source, options, runtime, root)
    # Validate topology and point offsets before any paid query.
    calculate(data, [0] * len(data['structures']), float(options.get('tolerance') or 2))
    if options.get('elevationSource') == 'arquivo':
        from geo import number
        try:
            elevations = [number(item['cota_terreno']) for item in data['structures']]
        except (ValueError, KeyError):
            raise ValueError('Todas as estruturas precisam de cota_terreno explícita no arquivo.') from None
        data['fonte_cota'] = 'cota_terreno informada no arquivo de entrada'
    else:
        transform = Transformer.from_crs(data['epsg'], 4326, always_xy=True)
        points = [transform.transform(*item['xy'], errcheck=True) for item in data['structures']]
        elevations = [row['cota_terreno_m'] for row in client.elevations(points)]
    result = calculate(data, elevations, float(options.get('tolerance') or 2))
    output.mkdir(parents=True, exist_ok=False)
    outputs = export_tables(result, output)
    target = output / ('estudo.kmz' if Path(source).suffix.lower() in {'.kmz', '.kml'} else 'estudo.dwg')
    if target.suffix == '.kmz':
        write_kmz(data, result, target)
    else:
        write_dwg(data, result, target, runtime, root)
    print(f"{len(result['structures'])} estruturas; {len(result['segments'])} trechos; extensão analisada {sum(t['extensao_m'] for t in result['segments']):.3f} m. Declividade do terreno.", flush=True)
    return [*outputs, target]

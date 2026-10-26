"""Coordenadas e consultas Google genéricas, sem dados de contratos."""
from __future__ import annotations

import csv
import json
import math
import os
from pathlib import Path
import socket
import urllib.error
import urllib.parse
import urllib.request

from pyproj import Transformer


class GoogleError(RuntimeError):
    pass


def google_key(env_file=None):
    key = os.environ.get('GOOGLE_MAPS_API_KEY', '')
    if not key and env_file:
        for line in Path(env_file).read_text(encoding='utf-8-sig').splitlines():
            name, sep, value = line.partition('=')
            if sep and name.strip() == 'GOOGLE_MAPS_API_KEY':
                key = value.strip().strip('\"\'')
                break
    return key


def number(value) -> float:
    result = float(str(value).strip().replace(',', '.'))
    if not math.isfinite(result):
        raise ValueError('Coordenada não finita.')
    return result


def lonlat(x, y):
    if not (-180 <= x <= 180 and -90 <= y <= 90):
        raise ValueError('Longitude/latitude fora do intervalo permitido.')
    return x, y


def request_json(url):
    with urllib.request.urlopen(url, timeout=30) as response:
        return json.load(response)


class GoogleClient:
    def __init__(self, key: str, request=request_json):
        if not key:
            raise GoogleError('Defina GOOGLE_MAPS_API_KEY ou escolha seu arquivo de configuração.')
        self.key, self.request = key, request

    def call(self, endpoint, params):
        url = f'https://maps.googleapis.com/maps/api/{endpoint}/json?' + urllib.parse.urlencode({**params, 'key': self.key})
        try:
            result = self.request(url)
        except (TimeoutError, socket.timeout):
            raise GoogleError('Google: tempo limite da consulta excedido. Tente novamente mais tarde.') from None
        except (urllib.error.URLError, OSError, ValueError):
            raise GoogleError('Não foi possível consultar o Google. Confira a conexão e a disponibilidade do serviço.') from None
        status = result.get('status')
        if status not in {'OK', 'ZERO_RESULTS'}:
            # Do not propagate URLs, API keys or provider-controlled error_message.
            raise GoogleError(f'Google: {status or "RESPOSTA_INVALIDA"}. Confira permissões da API, chave e quota.')
        return result

    def candidates(self, *, address=None, x=None, y=None):
        params = {'language': 'pt-BR'}
        if address is not None:
            if not str(address).strip():
                raise ValueError('Endereço vazio.')
            params['address'] = address
        else:
            lonlat(x, y)
            params['latlng'] = f'{y},{x}'
        result = self.call('geocode', params)
        return [{'address': item.get('formatted_address', ''),
                 'latitude': item['geometry']['location']['lat'],
                 'longitude': item['geometry']['location']['lng'],
                 'precision': item['geometry'].get('location_type', ''),
                 'partial': bool(item.get('partial_match', False))}
                for item in result.get('results', [])]

    def geocode(self, *, address=None, x=None, y=None):
        params = {'language': 'pt-BR'}
        if address is not None:
            if not str(address).strip():
                raise ValueError('Endereço vazio.')
            params['address'] = address
        else:
            lonlat(x, y)
            params['latlng'] = f'{y},{x}'
        result = self.call('geocode', params)
        if not result.get('results'):
            return {'status_resultado': 'SEM_RESULTADO', 'mensagem_resultado': 'Google não encontrou resultado.'}
        item = result['results'][0]
        location = item['geometry']['location']
        return {'endereco_resultado': item.get('formatted_address', ''),
                'longitude_resultado': location['lng'], 'latitude_resultado': location['lat'],
                'precisao_resultado': item['geometry'].get('location_type', ''),
                'correspondencia_parcial': bool(item.get('partial_match', False)),
                'status_resultado': 'REVISAR' if item.get('partial_match') else 'OK',
                'mensagem_resultado': 'Resultado Google; confira a localização antes de adotar.'}

    def elevations(self, points):
        outputs = []
        for i in range(0, len(points), 100):
            batch = points[i:i + 100]
            for xy in batch:
                lonlat(*xy)
            locations = '|'.join(f'{y},{x}' for x, y in batch)
            result = self.call('elevation', {'locations': locations})
            if len(result.get('results', [])) != len(batch):
                raise GoogleError('Google Elevation retornou uma resposta incompleta.')
            for item in result['results']:
                elevation = number(item['elevation'])
                outputs.append({'cota_terreno_m': elevation, 'resolucao_m': item.get('resolution', ''),
                                'fonte_cota': 'Google Elevation — terreno estimado', 'status_resultado': 'OK',
                                'mensagem_resultado': 'Resolução da grade não é garantia de precisão altimétrica.'})
        return outputs


def read_table(source):
    with Path(source).open(encoding='utf-8-sig', newline='') as stream:
        sample = stream.read(8192)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(stream, dialect=dialect)
        columns = reader.fieldnames or []
        if not columns or len(set(columns)) != len(columns) or any(not col.strip() for col in columns):
            raise ValueError('CSV deve ter cabeçalhos únicos e não vazios, em UTF-8.')
        rows = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError('CSV com quantidade de colunas inconsistente.')
            rows.append(row)
            if len(rows) > 100000:
                raise ValueError('Limite de 100.000 linhas por arquivo.')
        if not rows:
            raise ValueError('CSV sem linhas de dados.')
        return columns, rows


def inspect_table(source, tool):
    columns, rows = read_table(source)
    return {'columns': columns, 'rows': len(rows),
            'requests': math.ceil(len(rows) / 100) if tool == 'altitude' else len(rows) if tool in {'geocodificar', 'enderecos'} else 0}


def write_tables(rows, folder, stem='resultado'):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    columns = list(dict.fromkeys(key for row in rows for key in row))
    csv_path, xlsx_path = Path(folder) / f'{stem}.csv', Path(folder) / f'{stem}.xlsx'
    with csv_path.open('x', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, delimiter=';')
        writer.writeheader()
        for row in rows:
            # Neutralize spreadsheet formulas in untrusted textual input.
            safe = {key: "'" + value if isinstance(value, str) and value.startswith(('=', '+', '@')) else value for key, value in row.items()}
            writer.writerow(safe)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Resultados'
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(col, '') for col in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24
    book.save(xlsx_path)
    return [csv_path, xlsx_path]


def table(tool, source, output, options, client=None):
    output = Path(output)
    if output.exists():
        raise FileExistsError('A pasta de saída já existe.')
    columns, rows = read_table(source)
    x_col, y_col = options.get('xColumn', 'x'), options.get('yColumn', 'y')
    address_col = options.get('addressColumn', 'endereco')
    needed = [address_col] if tool == 'geocodificar' else [x_col, y_col]
    if any(col not in columns for col in needed):
        raise ValueError(f'Colunas não encontradas. Disponíveis: {", ".join(columns)}')
    transform = None
    if tool != 'geocodificar':
        source_crs = options.get('sourceCrs') or '4326'
        target = (options.get('targetCrs') or '4326') if tool == 'converter' else '4326'
        transform = Transformer.from_crs(source_crs, target, always_xy=True)
    result, elevations, positions = [], [], []
    stopped = ''
    for row in rows:
        item = {**row}
        if stopped:
            item.update(status_resultado='NAO_CONSULTADO', mensagem_resultado='Consultas interrompidas após falha do Google; nenhum novo pedido enviado.')
            result.append(item)
            continue
        try:
            if tool == 'geocodificar':
                item.update(client.geocode(address=row[address_col]))
            else:
                xy = transform.transform(number(row[x_col]), number(row[y_col]), errcheck=True)
                if tool == 'converter':
                    item.update(x_resultado=xy[0], y_resultado=xy[1], status_resultado='OK', mensagem_resultado='')
                elif tool == 'enderecos':
                    item.update(client.geocode(x=xy[0], y=xy[1]))
                elif tool == 'altitude':
                    lonlat(*xy)
                    positions.append(len(result)); elevations.append(xy)
                else:
                    raise ValueError('Ferramenta tabular desconhecida.')
        except GoogleError as error:
            stopped = str(error)
            item.update(status_resultado='ERRO_GOOGLE', mensagem_resultado=stopped)
        except (ValueError, TypeError, KeyError, RuntimeError) as error:
            item.update(status_resultado='ERRO', mensagem_resultado=str(error))
        result.append(item)
    if elevations:
        for offset in range(0, len(elevations), 100):
            indices = positions[offset:offset + 100]
            if stopped:
                for index in indices:
                    result[index].update(status_resultado='NAO_CONSULTADO', mensagem_resultado='Consulta interrompida após falha do Google.')
                continue
            try:
                batch = client.elevations(elevations[offset:offset + 100])
                for index, value in zip(indices, batch, strict=True):
                    result[index].update(value)
            except GoogleError as error:
                stopped = str(error)
                for index in indices:
                    result[index].update(status_resultado='ERRO_GOOGLE', mensagem_resultado=stopped)
    output.mkdir(parents=True, exist_ok=False)
    return write_tables(result, output)

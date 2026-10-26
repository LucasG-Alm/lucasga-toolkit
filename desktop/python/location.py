"""Consultas individuais, lotes e intercâmbio KML/KMZ do painel de localização."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path, PurePosixPath
import posixpath
import re
from urllib.parse import unquote, urlsplit
import uuid
import xml.etree.ElementTree as ET
import zipfile

from pyproj import CRS, Transformer
from pyproj.crs import ProjectedCRS
from pyproj.crs.coordinate_operation import UTMConversion

import geo

NS = 'http://www.opengis.net/kml/2.2'
ET.register_namespace('', NS)
def tag(name):
    return f'{{{NS}}}{name}'


def angle(value, latitude=False):
    text = str(value).strip().upper().replace(',', '.')
    # The exponent E belongs to a decimal number, before interpreting E as east.
    if re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:E[+-]?\d+)?', text):
        return geo.number(text)
    letters = re.findall('[NSEWO]', text)
    if len(letters) > 1 or (letters and letters[0] not in ('NS' if latitude else 'EWO')):
        raise ValueError('Hemisfério incompatível com a coordenada.')
    hemisphere = letters[0] if letters else ''
    text = re.sub('[NSEWO]', '', text).strip()
    pattern = r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)'
    parts = re.findall(pattern, text)
    residual = re.sub(pattern, '', text)
    if not 1 <= len(parts) <= 3 or re.sub(r'''[\s°º'"′″:]''', '', residual):
        raise ValueError('Coordenada inválida; use graus decimais ou graus, minutos e segundos.')
    values = [float(part) for part in parts]
    if any(x < 0 or x >= 60 for x in values[1:]):
        raise ValueError('Minutos e segundos devem estar entre 0 e 60.')
    negative = parts[0].startswith('-')
    if hemisphere and ((negative and hemisphere in 'NE') or (parts[0].startswith('+') and hemisphere in 'SWO')):
        raise ValueError('Sinal e hemisfério contraditórios.')
    result = abs(values[0]) + sum(v / (60 ** i) for i, v in enumerate(values[1:], 1))
    return -result if negative or hemisphere in ('S', 'W', 'O') else result


def parse_coordinate_text(text):
    """Separa coordenadas coladas; a validação de faixa UTM ocorre em normalize."""
    if not isinstance(text, str) or not text.strip() or len(text) > 4000:
        return None
    text = text.strip()
    try:
        if re.search(r'\b[XYZ]\s*=', text, re.I):
            text = re.sub(r'^at\s+point\s*,?\s*', '', text, flags=re.I)
            number = r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?'
            pattern = rf'\b([XYZ])\s*=\s*({number})(?=\s|[,;]|$)'
            matches = list(re.finditer(pattern, text, re.I))
            if re.sub(r'[\s,;]', '', re.sub(pattern, '', text, flags=re.I)):
                return None
            axes = [m[1].upper() for m in matches]
            if len(axes) != len(set(axes)) or not {'X', 'Y'} <= set(axes):
                return None
            values = {m[1].upper(): geo.number(m[2]) for m in matches}
            return {'kind': 'xy', 'easting': values['X'], 'northing': values['Y'],
                    **({'z': values['Z']} if 'Z' in values else {})}
        if ';' in text:
            parts = text.split(';')
        elif text.count(',') == 1:
            # A bare decimal comma (23,45) is ambiguous with a Lat/Lon pair.
            if not re.search(r'''[.°º′″'"NSEWO]|,\s|\s,|,[+-]''', text, re.I):
                return None
            parts = text.split(',')
        elif ',' not in text:
            hemispheres = re.fullmatch(r'(.+?[NS])\s+(.+?[EWO])', text, re.I)
            parts = list(hemispheres.groups()) if hemispheres else text.split()
        else:
            return None
        if len(parts) != 2:
            return None
        latitude, longitude = angle(parts[0], True), angle(parts[1])
        geo.lonlat(longitude, latitude)
        return {'kind': 'latlon', 'latitude': latitude, 'longitude': longitude}
    except (ValueError, TypeError, OverflowError):
        return None


def utm_crs(zone, hemisphere, datum):
    if geo.number(zone) != int(geo.number(zone)):
        raise ValueError('Fuso UTM deve ser inteiro.')
    zone = int(zone)
    if not 1 <= zone <= 60 or hemisphere not in ('S', 'N') or datum not in ('SIRGAS2000', 'WGS84'):
        raise ValueError('Informe fuso UTM entre 1 e 60, hemisfério e datum válidos.')
    return ProjectedCRS(name=f'{datum} / UTM zone {zone}{hemisphere}',
                        conversion=UTMConversion(str(zone), hemisphere),
                        geodetic_crs=CRS.from_epsg(4674 if datum == 'SIRGAS2000' else 4326))


def normalize(data):
    mode = data.get('mode', 'latlon')
    datum = data.get('datum') or 'SIRGAS2000'
    if datum not in ('SIRGAS2000', 'WGS84'):
        raise ValueError('Datum inválido; escolha SIRGAS2000 ou WGS84.')
    if mode == 'latlon':
        y, x = angle(data.get('latitude', ''), True), angle(data.get('longitude', ''))
    elif mode in ('utm', 'epsg'):
        x, y = geo.number(data.get('easting', '')), geo.number(data.get('northing', ''))
        if mode == 'utm':
            if not data.get('zone') or not data.get('hemisphere'):
                raise ValueError('Informe fuso e hemisfério UTM.')
            source = utm_crs(data['zone'], data['hemisphere'], datum)
            if not 100000 <= x <= 900000 or not 0 <= y <= 10000000:
                raise ValueError('E/N fora do intervalo UTM.')
        else:
            if not data.get('sourceCrs'):
                raise ValueError('Informe o CRS de origem.')
            source = CRS.from_user_input(data['sourceCrs'])
        x, y = Transformer.from_crs(source, 4326, always_xy=True).transform(x, y, errcheck=True)
    else:
        raise ValueError('Modo de coordenadas inválido.')
    geo.lonlat(x, y)
    if not math.isfinite(x) or not math.isfinite(y):
        raise ValueError('Coordenada não finita.')
    point = {'id': str(uuid.uuid4()), 'name': 'Ponto', 'latitude': y, 'longitude': x,
             'datum': datum, 'source': data.get('origin') or 'user'}
    zone = min(60, int((x + 180) // 6) + 1)
    if 56 <= y < 64 and 3 <= x < 12:
        zone = 32
    if 72 <= y < 84 and 0 <= x < 42:
        zone = 31 if x < 9 else 33 if x < 21 else 35 if x < 33 else 37
    hemisphere = 'S' if y < 0 else 'N'
    if data.get('targetZone') is not None or data.get('targetHemisphere') is not None:
        zone, hemisphere = data.get('targetZone'), data.get('targetHemisphere')
        target = utm_crs(zone, hemisphere, datum)
        if (y < 0 and hemisphere != 'S') or (y > 0 and hemisphere != 'N'):
            raise ValueError('Hemisfério UTM incompatível com a latitude.')
    else:
        target = utm_crs(zone, hemisphere, datum)
    if -80 <= y <= 84:
        easting, northing = Transformer.from_crs(4326, target, always_xy=True).transform(x, y, errcheck=True)
        point.update(zone=zone, hemisphere=hemisphere, easting=easting, northing=northing)
    else:
        point['message'] = 'Ponto válido fora da cobertura UTM (-80 a 84 graus).'
    if data.get('targetCrs'):
        tx, ty = Transformer.from_crs(4326, data['targetCrs'], always_xy=True).transform(x, y, errcheck=True)
        point.update(targetX=tx, targetY=ty, targetCrs=data['targetCrs'])
    return point


def fatal(error):
    return any(word in str(error) for word in ('REQUEST_DENIED', 'OVER_QUERY_LIMIT', 'OVER_DAILY_LIMIT', 'API_KEY', 'chave Google ausente', 'Defina GOOGLE_MAPS_API_KEY'))


def query(data, services, client=None):
    forward = data.get('mode') == 'address'
    address, elevation = bool(services.get('address')), bool(services.get('elevation'))
    if (forward or address or elevation) and not services.get('acknowledged'):
        raise ValueError('Confirme a previsão de consultas Google antes de consultar.')
    result = {'requests': {'geocode': 0, 'elevation': 0}, 'errors': {}, 'haltGoogle': False}
    if not forward:
        result['point'] = normalize(data)
    if forward or address:
        try:
            if client is None:
                raise geo.GoogleError('Defina GOOGLE_MAPS_API_KEY ou escolha seu arquivo de configuração.')
            result['requests']['geocode'] = 1
            kwargs = {'address': data.get('address', '')} if forward else {'x': result['point']['longitude'], 'y': result['point']['latitude']}
            candidates = client.candidates(**kwargs)
            if forward:
                result['candidates'] = candidates
                if len(candidates) == 1:
                    c = candidates[0]
                    result['point'] = normalize({**data, 'mode': 'latlon', 'latitude': str(c['latitude']), 'longitude': str(c['longitude']), 'origin': 'google'})
                    result['address'] = c['address']
            elif candidates:
                result['address'] = candidates[0]['address']
            if not candidates:
                result['errors']['address'] = 'Google não encontrou resultado.'
        except geo.GoogleError as error:
            result['errors']['address'] = str(error)
            result['haltGoogle'] = fatal(error)
    if elevation and 'point' in result:
        if result['haltGoogle']:
            result['errors']['elevation'] = 'Consulta não enviada após falha de chave ou quota.'
        else:
            try:
                if client is None:
                    raise geo.GoogleError('Defina GOOGLE_MAPS_API_KEY ou escolha seu arquivo de configuração.')
                result['requests']['elevation'] = 1
                p = result['point']
                item = client.elevations([(p['longitude'], p['latitude'])])[0]
                result['elevation'] = {'meters': geo.number(item['cota_terreno_m']), 'source': item['fonte_cota']}
                if item.get('resolucao_m') not in ('', None):
                    result['elevation']['resolution'] = geo.number(item['resolucao_m'])
            except geo.GoogleError as error:
                result['errors']['elevation'] = str(error)
                result['haltGoogle'] = result['haltGoogle'] or fatal(error)
    if 'point' in result:
        if result.get('address'):
            result['point']['address'] = result['address']
        if result.get('elevation'):
            result['point']['elevation'] = result['elevation']['meters']
            result['point']['elevationSource'] = result['elevation']['source']
            if 'resolution' in result['elevation']:
                result['point']['elevationResolution'] = result['elevation']['resolution']
    return result


def xml_root(data):
    if len(data) > 32 * 1024 * 1024 or b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('KML excede o limite ou contém entidades XML não permitidas.')
    return ET.fromstring(data)


def kml_document(source):
    source = Path(source)
    if source.suffix.lower() == '.kml':
        return xml_root(source.read_bytes()), {}
    assets = {}
    with zipfile.ZipFile(source) as archive:
        infos = archive.infolist()
        if len(infos) > 2048 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
            raise ValueError('KMZ excede o limite de segurança.')
        names = set()
        for item in infos:
            name = item.filename
            path = PurePosixPath(name)
            if path.is_absolute() or '..' in path.parts or '\\' in name or ':' in name or name in names:
                raise ValueError('KMZ contém caminho inseguro ou duplicado.')
            names.add(name)
        kmls = [i.filename for i in infos if i.filename.lower().endswith('.kml')]
        if 'doc.kml' in kmls:
            main = 'doc.kml'
        elif len(kmls) == 1:
            main = kmls[0]
        else:
            raise ValueError('KMZ com múltiplos KML deve identificar o principal como doc.kml na raiz.')
        root = xml_root(archive.read(main))
        assets = {i.filename: archive.read(i) for i in infos if i.filename != main and not i.is_dir()}
        if PurePosixPath(main).parent != PurePosixPath('.'):
            for node in [*root.iter(tag('href')), *root.iter(tag('styleUrl'))]:
                href = (node.text or '').strip()
                parsed = urlsplit(href)
                if href and not parsed.scheme and not parsed.netloc and not href.startswith('#'):
                    relative = posixpath.normpath(str(PurePosixPath(main).parent / unquote(parsed.path)))
                    if relative.startswith('../') or relative.startswith('/') or '\\' in relative:
                        raise ValueError('KMZ contém referência insegura.')
                    node.text = relative + (f'#{parsed.fragment}' if parsed.fragment else '')
    return root, assets


def kml_points(source):
    root, _ = kml_document(source)
    points = []
    for i, mark in enumerate(root.iter(tag('Placemark'))):
        nodes = list(mark.iter(tag('Point')))
        for j, node in enumerate(nodes):
            coords = node.find(tag('coordinates'))
            if coords is None or not coords.text:
                continue
            xy = coords.text.strip().split()[0].split(',')
            x, y = geo.lonlat(geo.number(xy[0]), geo.number(xy[1]))
            name = mark.findtext(tag('name')) or f'Ponto {i + 1}'
            points.append({'id': f'kml-{i}-{j}', 'name': name, 'longitude': x, 'latitude': y, 'source': 'imported'})
    return points


def read_input(source, options, inspection=False):
    source = Path(source)
    if source.suffix.lower() in ('.kml', '.kmz'):
        points = kml_points(source)
        ids = options.get('pointIds')
        selected = set(json.loads(ids)) if isinstance(ids, str) and ids else set(ids or [])
        rows = [p for p in points if p['id'] in selected]
        return ['id', 'name', 'longitude', 'latitude'], rows, {'points': [{'id': p['id'], 'name': p['name']} for p in points], 'rows': len(points)}
    if source.suffix.lower() == '.csv':
        cols, rows = geo.read_table(source)
        return cols, rows, {}
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Selecione CSV, XLSX, KML ou KMZ.')
    from openpyxl import load_workbook
    with source.open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            sheets = book.sheetnames
            sheet = options.get('sheet')
            if not sheet and len(sheets) > 1:
                if inspection:
                    return [], [], {'sheets': sheets, 'note': 'Selecione uma aba para inspecionar os dados.'}
                raise ValueError('Selecione explicitamente uma aba da planilha.')
            if sheet and sheet not in sheets:
                raise ValueError('A aba selecionada não existe.')
            iterator = book[sheet or sheets[0]].iter_rows(values_only=True)
            cols = [str(v).strip() if v is not None else '' for v in next(iterator, ())]
            if not cols or any(not c for c in cols) or len(cols) != len(set(cols)):
                raise ValueError('Planilha deve ter cabeçalhos únicos e não vazios.')
            rows = []
            for row in iterator:
                if any(v is not None for v in row):
                    rows.append(dict(zip(cols, [v if v is not None else '' for v in row])))
                if len(rows) > 100000:
                    raise ValueError('Limite de 100.000 linhas por arquivo.')
            return cols, rows, {'sheets': sheets}
        finally:
            book.close()


def row_input(row, options, kml=False):
    mode = 'latlon' if kml else options.get('mode') or 'latlon'
    x = row.get('longitude') if kml else row.get(options.get('xColumn') or 'x', '')
    y = row.get('latitude') if kml else row.get(options.get('yColumn') or 'y', '')
    return {**options, 'mode': mode, 'latitude': str(y), 'longitude': str(x), 'easting': str(x), 'northing': str(y),
            'address': row.get(options.get('addressColumn') or 'endereco', ''), 'origin': 'imported'}


def inspect(source, options):
    cols, rows, extra = read_input(source, options, True)
    kml = Path(source).suffix.lower() in ('.kml', '.kmz')
    valid = 0
    for row in rows:
        try:
            data = row_input(row, options, kml)
            if data['mode'] == 'address':
                if not str(data['address']).strip():
                    raise ValueError('Endereço vazio.')
            else:
                normalize(data)
            valid += 1
        except (ValueError, TypeError, RuntimeError):
            pass
    geocode = valid if options.get('mode') == 'address' or options.get('lookupAddress') else 0
    elevation = math.ceil(valid / 100) if options.get('lookupElevation') else 0
    return {'columns': cols, 'rows': len(rows), **extra, 'validRows': valid,
            'requests': geocode + elevation, 'services': {'geocode': geocode, 'elevation': elevation}}


def execute(source, output, options, client=None):
    output = Path(output)
    if output.exists():
        raise FileExistsError('A pasta de saída já existe; escolha uma nova.')
    cols, rows, _ = read_input(source, options)
    kml = Path(source).suffix.lower() in ('.kml', '.kmz')
    if not rows:
        raise ValueError('Selecione ao menos um ponto ou uma planilha com dados.')
    address = options.get('mode') == 'address' or options.get('lookupAddress')
    elevation = options.get('lookupElevation')
    if (address or elevation) and not options.get('googleAcknowledged'):
        raise ValueError('Confirme a previsão de consultas Google.')
    results, elevation_items = [], []
    stopped = False
    for index, row in enumerate(rows):
        item = {**row, 'status_endereco': 'NAO_SOLICITADO', 'status_altitude': 'NAO_SOLICITADO'}
        try:
            data = row_input(row, options, kml)
            if stopped and data['mode'] == 'address':
                item.update(status_resultado='NAO_CONSULTADO', status_endereco='NAO_CONSULTADO')
                results.append(item); continue
            result = query(data, {'address': bool(address) and not stopped, 'elevation': False, 'acknowledged': True}, client)
            stopped = stopped or result['haltGoogle']
            point = result.get('point')
            if address:
                item['status_endereco'] = 'NAO_CONSULTADO' if stopped and not result['errors'] else 'ERRO_GOOGLE' if result['errors'].get('address') else 'OK' if result.get('address') else 'AMBIGUO'
            if result['errors']:
                item['mensagem_resultado'] = '; '.join(result['errors'].values())
            if point:
                item.update(longitude_resultado=point['longitude'], latitude_resultado=point['latitude'],
                            easting_resultado=point.get('easting', ''), northing_resultado=point.get('northing', ''),
                            fuso_resultado=point.get('zone', ''), hemisferio_resultado=point.get('hemisphere', ''), datum_resultado=point['datum'],
                            origem_resultado=point['source'], id_resultado=str(row.get(options.get('idColumn') or 'id') or index + 1),
                            nome_resultado=str(row.get(options.get('nameColumn') or 'name') or f'Ponto {index + 1}'),
                            status_resultado='OK' if not result['errors'] and item['status_endereco'] != 'NAO_CONSULTADO' else 'PARCIAL')
                if result.get('address'):
                    item['endereco_resultado'] = result['address']
                if point.get('targetCrs'):
                    item.update(x_resultado=point['targetX'], y_resultado=point['targetY'], crs_resultado=point['targetCrs'])
                if elevation:
                    elevation_items.append((len(results), (point['longitude'], point['latitude'])))
            else:
                item['status_resultado'] = 'ERRO_GOOGLE' if result['errors'] else 'AMBIGUO'
                item['candidatos_resultado'] = json.dumps(result.get('candidates', []), ensure_ascii=False)
        except (ValueError, TypeError, RuntimeError) as error:
            item.update(status_resultado='ERRO', mensagem_resultado=str(error))
        results.append(item)
    for offset in range(0, len(elevation_items), 100):
        batch = elevation_items[offset:offset + 100]
        if stopped:
            for index, _ in batch:
                results[index].update(status_altitude='NAO_CONSULTADO', status_resultado='PARCIAL')
            continue
        try:
            if client is None:
                raise geo.GoogleError('Defina GOOGLE_MAPS_API_KEY ou escolha seu arquivo de configuração.')
            values = client.elevations([xy for _, xy in batch])
            for (index, _), value in zip(batch, values, strict=True):
                results[index].update(cota_terreno_m=value['cota_terreno_m'], resolucao_m=value.get('resolucao_m', ''), fonte_cota=value['fonte_cota'], status_altitude='OK')
        except geo.GoogleError as error:
            stopped = fatal(error)
            for index, _ in batch:
                results[index].update(status_altitude='ERRO_GOOGLE', status_resultado='PARCIAL', mensagem_altitude=str(error))
    output.mkdir(parents=True, exist_ok=False)
    outputs = geo.write_tables(results, output)
    attention = sum(r['status_resultado'] != 'OK' for r in results)
    return {'outputs': outputs, 'pause': stopped, 'warning': f'{attention} linha(s) requerem revisão.' if attention else ''}


def read_results(source):
    cols, rows, _ = read_input(source, {'sheet': 'Resultados'})
    points = []
    for i, row in enumerate(rows):
        try:
            x, y = geo.lonlat(geo.number(row['longitude_resultado']), geo.number(row['latitude_resultado']))
            point = {'id': f"result-{i + 1}-{row.get('id_resultado') or ''}", 'name': str(row.get('nome_resultado') or f'Ponto {i + 1}'),
                     'longitude': x, 'latitude': y, 'source': row.get('origem_resultado') or 'imported', 'status': row.get('status_resultado', '')}
            if row.get('endereco_resultado'):
                point['address'] = str(row['endereco_resultado'])
            if row.get('cota_terreno_m') not in ('', None):
                point['elevation'] = geo.number(row['cota_terreno_m'])
            points.append(point)
        except (ValueError, TypeError, KeyError):
            continue
    return points


def export_points(points, target, base=None, crs=None):
    import point_exports
    target = Path(target)
    if target.suffix.lower() in ('.xlsx', '.dxf'):
        if base:
            raise ValueError('XLSX e DXF não aceitam arquivo-base.')
        return point_exports.export_points(points, target, crs)
    # Legacy KML callers may omit the name; preserve their existing default.
    if isinstance(points, list):
        points = [{**point, 'name': point.get('name', 'Ponto')} if isinstance(point, dict) else point
                  for point in points]
    points = point_exports.validate_points(points)
    if target.exists():
        raise FileExistsError('O destino já existe; escolha um arquivo novo.')
    if target.suffix.lower() not in ('.kml', '.kmz') or not points:
        raise ValueError('Selecione pontos e destino KML ou KMZ.')
    assets = {}
    if base:
        root, assets = kml_document(base)
        if Path(base).suffix.lower() == '.kmz' and target.suffix.lower() == '.kml' and assets:
            raise ValueError('Exporte como KMZ para preservar os recursos da base KMZ.')
        if Path(base).suffix.lower() == '.kml':
            for node in root.iter(tag('href')):
                href = (node.text or '').strip()
                if href and '://' not in href and not href.startswith('#'):
                    if target.suffix.lower() == '.kml':
                        if target.resolve().parent != Path(base).resolve().parent:
                            raise ValueError('Salve o KML na mesma pasta da base ou exporte KMZ para preservar recursos relativos.')
                        continue
                    path = unquote(urlsplit(href).path)
                    relative = PurePosixPath(path)
                    if relative.is_absolute() or '..' in relative.parts or '\\' in path or ':' in path:
                        raise ValueError('A base KML contém recurso relativo inseguro.')
                    resource = (Path(base).resolve().parent / path).resolve(strict=True)
                    if not resource.is_relative_to(Path(base).resolve().parent) or not resource.is_file():
                        raise ValueError('O recurso KML deve ficar dentro da pasta da base.')
                    if resource.suffix.lower() in ('.kml', '.kmz'):
                        raise ValueError('Bases KML com referências a outros KML/KMZ devem ser consolidadas antes da exportação.')
                    if resource.stat().st_size > 32 * 1024 * 1024:
                        raise ValueError('Recurso KML excede 32 MB.')
                    assets[path] = resource.read_bytes()
                    if len(assets) > 2047 or sum(len(content) for content in assets.values()) > 100 * 1024 * 1024:
                        raise ValueError('Recursos KML excedem o limite de segurança.')
    else:
        root = ET.Element(tag('kml'))
    document = root.find(tag('Document'))
    if document is None:
        original = list(root)
        for element in original:
            root.remove(element)
        document = ET.SubElement(root, tag('Document'))
        for element in original:
            document.append(element)
    folder = ET.SubElement(document, tag('Folder'))
    ET.SubElement(folder, tag('name')).text = 'Pontos do Toolkit'
    for point in points:
        x, y = geo.lonlat(geo.number(point['longitude']), geo.number(point['latitude']))
        mark = ET.SubElement(folder, tag('Placemark'))
        ET.SubElement(mark, tag('name')).text = str(point.get('name') or 'Ponto')
        description = str(point.get('description') or point.get('address') or '')
        if point.get('elevation') is not None:
            elevation = geo.number(point['elevation'])
            source = point.get('elevationSource') or 'não informada'
            description += ('\n' if description else '') + f'Cota terreno: {elevation!r} m. Fonte: {source}; resolução não garante precisão altimétrica.'
        if description:
            ET.SubElement(mark, tag('description')).text = description
        node = ET.SubElement(mark, tag('Point'))
        ET.SubElement(node, tag('altitudeMode')).text = 'clampToGround'
        ET.SubElement(node, tag('coordinates')).text = f'{x!r},{y!r}'
    data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    if target.suffix.lower() == '.kmz':
        with zipfile.ZipFile(target, 'x', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('doc.kml', data)
            for name, content in assets.items():
                archive.writestr(name, content)
    else:
        with target.open('xb') as stream:
            stream.write(data)
    return [target]

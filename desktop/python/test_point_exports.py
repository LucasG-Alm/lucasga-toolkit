from pathlib import Path
import sys

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import point_exports

CRS = {'zone': 23, 'hemisphere': 'S', 'datum': 'SIRGAS2000'}
POINT = {'id': 'p1', 'name': '=Árvore\nCentro', 'latitude': 0, 'longitude': -45,
         'source': 'user', 'elevation': 123.456789}


@pytest.mark.parametrize('latitude,longitude,zone', [(-.000001, -45, 23), (-23, -.000001, 30)])
def test_scientific_coordinates_reopen_exports_against_direct_projection(tmp_path, monkeypatch, latitude, longitude, zone):
    from openpyxl import load_workbook
    from pyproj import Transformer
    import ezdxf
    import geo
    def forbidden(*args, **kwargs):
        raise AssertionError('Exportação local não consulta Google')
    monkeypatch.setattr(geo, 'GoogleClient', forbidden)
    crs = {'zone': zone, 'hemisphere': 'S', 'datum': 'WGS84'}
    expected = Transformer.from_crs(4326, 32700 + zone, always_xy=True).transform(longitude, latitude)
    for extension in ['xlsx', 'dxf']:
        target = tmp_path / f'near-zero.{extension}'
        point_exports.export_points([{**POINT, 'latitude': latitude, 'longitude': longitude}], target, crs)
        if extension == 'xlsx':
            sheet = load_workbook(target).active
            actual = sheet['E2'].value, sheet['F2'].value
            assert sheet['C2'].value == latitude and sheet['D2'].value == longitude
        else:
            vector = list(ezdxf.readfile(target).modelspace().query('POINT'))[0].dxf.location
            actual = vector.x, vector.y
        assert actual == pytest.approx(expected, abs=1e-8)


@pytest.mark.parametrize('field', ['id', 'name', 'address', 'description', 'elevationSource'])
@pytest.mark.parametrize('extension', ['kml', 'kmz', 'xlsx', 'dxf'])
def test_xml_illegal_control_is_rejected_before_file(tmp_path, field, extension):
    import location
    target = tmp_path / f'illegal.{extension}'
    with pytest.raises(ValueError, match='inválido'):
        location.export_points([{**POINT, field: 'Ponto\x01A'}], target, crs=CRS)
    assert not target.exists()


def test_valid_xml_whitespace_and_escaped_characters_reopen(tmp_path):
    import location
    import xml.etree.ElementTree as ET
    target = tmp_path / 'valid.kml'
    location.export_points([{**POINT, 'name': '<Ponto> & A\tB\nC\rD'}], target)
    assert ET.parse(target).find('.//' + location.tag('Placemark')) is not None


def test_reopen_xlsx_dxf_precision_strings_units_and_zero_z(tmp_path):
    from openpyxl import load_workbook
    import ezdxf
    xlsx, dxf = tmp_path / 'pontos.xlsx', tmp_path / 'pontos.dxf'
    point_exports.export_points([POINT], xlsx, CRS)
    point_exports.export_points([POINT], dxf, CRS)
    sheet = load_workbook(xlsx).active
    row = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert row['Nome'] == POINT['name'] and sheet['B2'].data_type == 's'
    assert row['Latitude WGS84'] == 0 and row['Longitude WGS84'] == -45
    assert abs(row['E UTM (m)'] - 500000) < .001
    assert row['Cota terreno (m)'] == POINT['elevation']
    assert row['Fonte da cota'] is None
    assert row['Datum UTM'] == CRS['datum']
    drawing = ezdxf.readfile(dxf)
    entity = list(drawing.modelspace().query('POINT'))[0]
    assert abs(entity.dxf.location.x - 500000) < .001
    assert entity.dxf.location.z == 0 and drawing.header['$INSUNITS'] == 6
    assert list(drawing.modelspace().query('TEXT'))[0].dxf.text == '=Árvore Centro'
    assert 'SIRGAS2000 / UTM 23S' in dxf.read_text(encoding='utf-8')


def test_xlsx_preserves_projection_precision_metadata_and_formula_like_text(tmp_path):
    import location
    from openpyxl import load_workbook
    target = tmp_path / 'precision.xlsx'
    point = {**POINT, 'latitude': -23.1234567890123, 'longitude': -46.1234567890123,
             'address': '+Rua sintética', 'elevationSource': '@Levantamento',
             'elevationResolution': 12.34567890123}
    points = [{**point, 'name': name} for name in ['=1+1', '+Ponto', '-Ponto', '@Ponto']]
    point_exports.export_points(points, target, CRS)
    expected = location.normalize({'mode': 'latlon', 'latitude': str(point['latitude']),
        'longitude': str(point['longitude']), 'targetZone': 23, 'targetHemisphere': 'S'})
    sheet = load_workbook(target).active
    for index in range(2, 6):
        assert sheet.cell(index, 2).data_type == 's'
        assert sheet.cell(index, 10).data_type == 's'
        assert sheet.cell(index, 12).data_type == 's'
        assert abs(sheet.cell(index, 5).value - expected['easting']) < 1e-8
        assert abs(sheet.cell(index, 6).value - expected['northing']) < 1e-8
        assert sheet.cell(index, 13).value == point['elevationResolution']


@pytest.mark.parametrize('extension', ['xlsx', 'dxf'])
def test_existing_destination_and_invalid_later_point_never_overwrite(tmp_path, extension):
    target = tmp_path / f'points.{extension}'
    target.write_bytes(b'original')
    with pytest.raises(FileExistsError):
        point_exports.export_points([POINT], target, CRS)
    assert target.read_bytes() == b'original'
    other = tmp_path / f'invalid.{extension}'
    with pytest.raises(ValueError, match='p2.*CRS.*KMZ'):
        point_exports.export_points([POINT, {**POINT, 'id': 'p2', 'latitude': 23}], other, CRS)
    assert not other.exists()


@pytest.mark.parametrize('changes', [{'latitude': float('nan')}, {'name': ''},
    {'elevation': float('inf')}, {'elevationResolution': -1}, {'address': 42}])
def test_invalid_metadata_before_file(tmp_path, changes):
    target = tmp_path / 'invalid.xlsx'
    with pytest.raises(ValueError):
        point_exports.export_points([{**POINT, **changes}], target, CRS)
    assert not target.exists()


def test_empty_invalid_crs_and_base_rejected(tmp_path):
    import location
    for points, crs in [([], CRS), ([POINT], None), ([POINT], {**CRS, 'zone': True})]:
        with pytest.raises(ValueError):
            point_exports.export_points(points, tmp_path / 'p.xlsx', crs)
    with pytest.raises(ValueError, match='base'):
        location.export_points([POINT], tmp_path / 'p.xlsx', 'base.kml', CRS)


@pytest.mark.parametrize('extension', ['xlsx', 'dxf'])
def test_exclusive_creation_race_preserves_competing_file(tmp_path, monkeypatch, extension):
    target = tmp_path / f'race.{extension}'
    original_open = Path.open
    def racing_open(path, mode='r', *args, **kwargs):
        if path == target and mode == 'xb':
            with original_open(path, 'wb') as competitor:
                competitor.write(b'competing file')
        return original_open(path, mode, *args, **kwargs)
    monkeypatch.setattr(Path, 'open', racing_open)
    with pytest.raises(FileExistsError):
        point_exports.export_points([POINT], target, CRS)
    assert target.read_bytes() == b'competing file'


def test_kml_elevation_source_is_explicit_and_never_assumed(tmp_path):
    import location
    target = tmp_path / 'unknown.kml'
    location.export_points([POINT], target)
    text = target.read_text(encoding='utf-8')
    assert 'Google' not in text and 'não informada' in text
    known = tmp_path / 'known.kml'
    location.export_points([{**POINT, 'elevationSource': 'Levantamento sintético'}], known)
    assert 'Levantamento sintético' in known.read_text(encoding='utf-8')


@pytest.mark.parametrize('extension,library', [('xlsx', 'openpyxl'), ('dxf', 'ezdxf')])
def test_missing_library_does_not_create_destination(tmp_path, monkeypatch, extension, library):
    import builtins
    original_import = builtins.__import__
    def without_library(name, *args, **kwargs):
        if name == library:
            raise ModuleNotFoundError(f'No module named {library}')
        return original_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', without_library)
    target = tmp_path / f'missing.{extension}'
    with pytest.raises(ModuleNotFoundError, match=library):
        point_exports.export_points([POINT], target, CRS)
    assert not target.exists()

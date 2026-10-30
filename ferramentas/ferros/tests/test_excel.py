"""Planilha FERRO gerada pelo app. O Excel real so com FERROS_EXCEL=1."""

import os
from pathlib import Path

import openpyxl
import pytest

from cotador.excel import PlanilhaAberta, gerar_planilha
from cotador.grade import Bloqueado, TextoLido, lista_da_grade, montar_grade
from cotador.resumo import PESO_POR_M, montar_resumo


def grade(*textos):
    return montar_grade([TextoLido(f"H{i}", t, "ADUELA", 0.0, 0.0) for i, t in enumerate(textos)])


def test_gera_ao_lado_do_dwg_com_uma_tag_por_linha(tmp_path):
    """R-FER-XLS-001 / R-FER-XLS-002"""
    g = grade("2x4 N11 Ø12.5 c=391", "3N2 %%C10 c=100")
    g.multiplicadores[11] = 6
    saida = gerar_planilha(g, PESO_POR_M, tmp_path / "prancha.dwg")
    assert saida == tmp_path / "prancha - FERRO.xlsx"
    ws = openpyxl.load_workbook(saida)["FERRO"]
    assert [ws["B4"].value, ws["C4"].value] == ["6x8N11 ∅12.5 c=391", "ADUELA"]
    assert ws["B5"].value == "3N2 ∅10 c=100"
    assert ws["B6"].value is None
    assert ws["F4"].value.startswith("=IFERROR(") and 'FIND("N",B4)' in ws["F4"].value


def test_formula_le_a_letra_da_posicao(tmp_path):
    """R-FER-XLS-002: grampo G nao e N."""
    ws = openpyxl.load_workbook(gerar_planilha(grade("4G3 Ø6.3 c=50"), PESO_POR_M, tmp_path / "p.dwg"))["FERRO"]
    assert 'FIND("G",B4)' in ws["D4"].value and 'FIND("G",B4)' in ws["F4"].value


def test_resumo_por_bitola_usa_o_intervalo_exato_e_o_peso_da_pagina(tmp_path):
    """R-FER-XLS-003"""
    g = grade("1N1 Ø10 c=100", "1N2 Ø8 c=100", "1N3 Ø10 c=100")
    ws = openpyxl.load_workbook(gerar_planilha(g, {**PESO_POR_M, 10.0: 0.63}, tmp_path / "p.dwg"))["FERRO"]
    linhas = {ws[f"E{r}"].value: r for r in range(7, 20) if ws[f"E{r}"].value is not None}
    assert set(linhas) >= {8.0, 10.0, "Total"}
    r10 = linhas[10.0]
    assert "$H$4:$H$6" in ws[f"F{r10}"].value and ws[f"G{r10}"].value == 0.63
    assert ws[f"I4"].value.startswith("=VLOOKUP(E4,")


def test_resumo_por_posicao_ao_lado(tmp_path):
    """R-FER-XLS-004"""
    ws = openpyxl.load_workbook(gerar_planilha(grade("2N11 Ø12.5 c=391", "1N2 Ø10 c=100"), PESO_POR_M, tmp_path / "p.dwg"))["FERRO"]
    assert ws["K8"].value == "N"
    posicoes = [ws[f"K{r}"].value for r in range(9, 13) if ws[f"K{r}"].value]
    assert posicoes == ["02", "11"]
    assert ws["M9"].value.startswith("=SUMIFS($F$4:$F$5,")


def test_antigo_vai_para_backup(tmp_path):
    """R-FER-XLS-001"""
    dwg = tmp_path / "p.dwg"
    gerar_planilha(grade("1N1 Ø10 c=100"), PESO_POR_M, dwg)
    gerar_planilha(grade("2N1 Ø10 c=100"), PESO_POR_M, dwg)
    backups = list((tmp_path / "_BACKUP").glob("p - FERRO (antes de *).xlsx"))
    assert len(backups) == 1
    assert openpyxl.load_workbook(backups[0])["FERRO"]["B4"].value == "1N1 ∅10 c=100"


def test_aberta_no_excel_nao_grava(tmp_path):
    """R-FER-XLS-001"""
    dwg = tmp_path / "p.dwg"
    gerar_planilha(grade("1N1 Ø10 c=100"), PESO_POR_M, dwg)
    (tmp_path / "~$p - FERRO.xlsx").write_bytes(b"")
    with pytest.raises(PlanilhaAberta, match="feche"):
        gerar_planilha(grade("2N1 Ø10 c=100"), PESO_POR_M, dwg)
    assert openpyxl.load_workbook(tmp_path / "p - FERRO.xlsx")["FERRO"]["B4"].value == "1N1 ∅10 c=100"


def test_sem_caminho_ou_com_vermelho_nao_gera(tmp_path):
    """R-FER-XLS-001 / R-FER-XLS-006"""
    with pytest.raises(ValueError, match="Salve o desenho"):
        gerar_planilha(grade("1N1 Ø10 c=100"), PESO_POR_M, None)
    with pytest.raises(Bloqueado):
        gerar_planilha(grade("3 N1 c/10"), PESO_POR_M, tmp_path / "p.dwg")
    assert not (tmp_path / "p - FERRO.xlsx").exists()


@pytest.mark.skipif(os.environ.get("FERROS_EXCEL") != "1", reason="FERROS_EXCEL=1 liga o Excel real")
def test_total_do_excel_e_o_da_pagina(tmp_path):
    """R-FER-XLS-005: o Excel recalcula as formulas e chega no mesmo total."""
    import win32com.client

    g = grade("2x4 N11 Ø12.5 c=391", "3N2 Ø10 c=100", "4G3 Ø6.3 c/15 c=50", "2x4 N11 Ø12.5 c=391")
    g.multiplicadores[11] = 6
    g.linhas[1].quantidade = 5
    _, esperado = montar_resumo(lista_da_grade(g), PESO_POR_M)
    saida = gerar_planilha(g, PESO_POR_M, tmp_path / "p.dwg")

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        livro = excel.Workbooks.Open(str(Path(saida)))
        livro.Application.CalculateFull()
        ws = livro.Worksheets("FERRO")
        total = next(ws.Cells(r, 8).Value for r in range(4, 40) if ws.Cells(r, 5).Value == "Total")
        livro.Close(False)
    finally:
        excel.Quit()
    assert round(total, 2) == esperado


EXCEL_2007 = {"TEXT", "MID", "FIND", "IFERROR", "VLOOKUP", "CONVERT", "SUMIFS", "SUM", "SUMPRODUCT", "MAX"}


def test_formulas_so_usam_funcoes_do_excel_2007(tmp_path):
    """R-FER-XLS-007: a empresa usa Excel 2007-2010."""
    import re

    g = grade("2x4 N11 Ø12.5 c=391", "3N2 Ø10 c=100", "4G3 Ø6.3 c=50")
    ws = openpyxl.load_workbook(gerar_planilha(g, PESO_POR_M, tmp_path / "p.dwg"))["FERRO"]
    usadas = {f for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")
              for f in re.findall(r"([A-Z][A-Z0-9._]*)\(", c.value)}
    assert usadas <= EXCEL_2007, usadas - EXCEL_2007


def test_total_em_negrito_italico_com_borda_dupla(tmp_path):
    """R-FER-XLS-008"""
    ws = openpyxl.load_workbook(gerar_planilha(grade("1N1 Ø10 c=100"), PESO_POR_M, tmp_path / "p.dwg"))["FERRO"]
    r = next(r for r in range(4, 20) if ws[f"E{r}"].value == "Total")
    for col in "EFGH":
        c = ws[f"{col}{r}"]
        assert c.font.b and c.font.i and c.border.top.style == "double", col
    assert ws["C3"].font.b and ws["E3"].font.b and ws["K7"].font.b

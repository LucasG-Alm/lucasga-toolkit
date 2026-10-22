"""Ortografia dos textos dos desenhos, sem LLM: dicionário Hunspell pt-BR + glossário.

revisar: lê os DWG da raiz da pasta, marca as palavras fora do dicionário e do glossário,
grava o relatório Excel e, para cada desenho com suspeita, uma cópia <nome>-REVISAO.dwg
com caixa e nota (layer _ANOTAÇÕES, cor 210) no papel, onde a caneta vermelha marcaria.
aprender: passa ao glossário as palavras marcadas na coluna Aceitar do relatório.
Regras em COMPORTAMENTO.md (R-ORT-*).
"""
from __future__ import annotations

import argparse
import datetime
import fnmatch
import math
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import ezdxf
import openpyxl
from ezdxf import bbox
from ezdxf.disassemble import recursive_decompose
from ezdxf.tools.text import plain_mtext

import layouts as lay

AQUI = Path(__file__).resolve().parent
DICIONARIO = AQUI / "dicionario" / "pt_BR"
GLOSSARIO = AQUI.parents[1] / "local" / "glossario.txt"     # fora do repositório público
ARMADILHAS = AQUI / "armadilhas.txt"                          # genérico, vai no repositório
IGNORAR_DECIMAL = AQUI / "ignorar_decimal.txt"                # layers fora da regra do decimal
CAMADA, COR = "_ANOTAÇÕES", 210
ACENTOS = {"a": "áàâã", "e": "éê", "i": "í", "o": "óôõ", "u": "úü", "c": "ç"}
LETRAS = "abcdefghijklmnopqrstuvwxyzáàâãéêíóôõúüç"
_PALAVRA = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+")


# ---------------------------------------------------------------- palavras
def palavras(texto: str) -> list[str]:
    """Sequências de letras; trecho (entre espaços) com dígito sai inteiro, letra solta também."""
    out = []
    for trecho in texto.split():
        if not any(ch.isdigit() for ch in trecho):
            out += [w for w in _PALAVRA.findall(trecho) if len(w) > 1]
    return out


@lru_cache(maxsize=2)
def _dicionario(caminho: str):
    from spylls.hunspell import Dictionary
    return Dictionary.from_files(caminho)


def _caixa_de(modelo: str, palavra: str) -> str:
    if modelo.isupper():
        return palavra.upper()
    return palavra.capitalize() if modelo[:1].isupper() else palavra


def _com_acento(w: str, trocas: int = 2) -> set[str]:
    w = w.lower()
    pos = [i for i, ch in enumerate(w) if ch in ACENTOS]
    vistos, fronteira = set(), {w}
    for _ in range(trocas):
        fronteira = {c[:i] + r + c[i + 1:] for c in fronteira for i in pos
                     if c[i] in ACENTOS for r in ACENTOS[c[i]]}
        vistos |= fronteira
    return vistos


def _uma_letra(w: str) -> set[str]:
    w = w.lower()
    cortes = [(w[:i], w[i:]) for i in range(len(w) + 1)]
    return ({a + b[1:] for a, b in cortes if b}
            | {a + b[1] + b[0] + b[2:] for a, b in cortes if len(b) > 1}
            | {a + c + b[1:] for a, b in cortes if b for c in LETRAS}
            | {a + c + b for a, b in cortes for c in LETRAS})


class Verificador:
    def __init__(self, glossario=(), dicionario: Path = DICIONARIO, armadilhas: dict | None = None,
                 ignorar_decimal: list | None = None):
        self.d = _dicionario(str(dicionario))
        self.glossario = {g.lower() for g in glossario}
        self.armadilhas = le_armadilhas(ARMADILHAS) if armadilhas is None else armadilhas
        self.ignorar_decimal = le_padroes(IGNORAR_DECIMAL) if ignorar_decimal is None else ignorar_decimal
        self._sug: dict[str, list[str]] = {}

    def suspeita(self, w: str) -> bool:
        if w.lower() in self.glossario:
            return False
        return w.lower() in self.armadilhas or not self.d.lookup(w)

    def sugestoes(self, w: str) -> list[str]:
        if w.lower() in self.armadilhas:
            return [_caixa_de(w, self.armadilhas[w.lower()])]
        if w not in self._sug:
            achou = [c for c in _com_acento(w) if self.d.lookup(c)]
            if not achou and len(w) <= 15:      # ~1 s por palavra; acima disso não compensa
                achou = [c for c in _uma_letra(w) if self.d.lookup(c)]
            self._sug[w] = sorted({_caixa_de(w, c) for c in achou})[:5]
        return self._sug[w]


def le_armadilhas(caminho: Path) -> dict[str, str]:
    """`errado = certo` por linha (# comenta): palavra válida sem acento que em projeto é erro."""
    caminho = Path(caminho)
    if not caminho.is_file():
        return {}
    out = {}
    for l in caminho.read_text(encoding="utf-8").splitlines():
        l = l.split("#", 1)[0]
        if "=" in l:
            errado, certo = (x.strip() for x in l.split("=", 1))
            if errado and certo:
                out[errado.lower()] = certo.lower()
    return out


def le_glossario(caminho: Path) -> set[str]:
    caminho = Path(caminho)
    if not caminho.is_file():
        return set()
    linhas = (l.split("#", 1)[0].strip() for l in caminho.read_text(encoding="utf-8").splitlines())
    return {l.lower() for l in linhas if l}


# ----------------------------------------- unidades, expressões e pontuação
_UNIDADES_SI = {
    "Km": "km", "KM": "km",
    "mts": "m", "Mts": "m", "MTS": "m", "mt": "m", "Mt": "m", "MT": "m",
    "mpa": "MPa", "Mpa": "MPa",
    "kpa": "kPa", "Kpa": "kPa", "KPA": "kPa",
    "kn": "kN", "KN": "kN",
    "kw": "kW", "KW": "kW",
    "kva": "kVA", "KVA": "kVA",
    "kv": "kV", "KV": "kV",
    "ka": "kA", "KA": "kA",
    "VAR": "var", "VAr": "var", "kVAr": "kvar", "KVAR": "kvar",
    "hz": "Hz", "HZ": "Hz", "Mhz": "MHz", "mhz": "MHz", "MHZ": "MHz",
    "kw/h": "kWh", "kW/h": "kWh", "KW/h": "kWh", "KW/H": "kWh", "kwh": "kWh", "KWH": "kWh",
    "kgs": "kg", "Kgs": "kg", "KGS": "kg", "Kg": "kg", "KG": "kg",
    "grs": "g", "Grs": "g", "GRS": "g", "gr": "g", "Gr": "g", "GR": "g",
    "hrs": "h", "Hrs": "h", "HRS": "h", "hr": "h", "Hr": "h", "HR": "h",
    "segs": "s", "Segs": "s", "SEGS": "s", "seg": "s", "Seg": "s", "SEG": "s",
    "l/seg": "L/s", "L/seg": "L/s", "l/SEG": "L/s", "L/SEG": "L/s",
    "m3/hr": "m³/h", "M3/HR": "m³/h", "m³/hr": "m³/h", "m3/h": "m³/h", "M3/H": "m³/h",
    "m3/dia": "m³/d", "m³/dia": "m³/d", "M3/DIA": "m³/d",
}
_RE_UNIDADE = re.compile(
    r"(?<![A-Za-zÀ-ÖØ-öø-ÿ\d])(\d+(?:[.,]\d+)?)\s*("
    + "|".join( sorted((re.escape(k) for k in _UNIDADES_SI), key=len, reverse=True) )
    + r")(?![A-Za-zÀ-ÖØ-öø-ÿ\d/])"
)

_EXPRESSOES: list[tuple[re.Pattern, object]] = [
    (re.compile(r"\bsub[-\s]+leito\b", re.I), "subleito"),
    (re.compile(r"\bsub[-\s]+solo\b", re.I), "subsolo"),
    (re.compile(r"\b(?:sub\s+base|subbase)\b", re.I), "sub-base"),
    (re.compile(r"\bcontra[-\s]+piso\b", re.I), "contrapiso"),
    (re.compile(r"\b(?:pre[-\s]+|pré\s+)moldad([oa]s?)\b", re.I), lambda m: f"pré-moldad{m.group(1).lower()}"),
    (re.compile(r"\b(?:pre[-\s]+|pré\s+)fabricad([oa]s?)\b", re.I), lambda m: f"pré-fabricad{m.group(1).lower()}"),
    (re.compile(r"\bguarda\s+corpos?\b", re.I), lambda m: "guarda-corpos" if m.group(0).lower().endswith("s") else "guarda-corpo"),
    (re.compile(r"\bguardas-corpos?\b", re.I), "guarda-corpos"),
    (re.compile(r"\b(?:p[eé]\s+direito|pe-direito)\b", re.I), "pé-direito"),
    (re.compile(r"\b(?:p[eé]-direitos|pés-direito)\b", re.I), "pés-direitos"),
    (re.compile(r"\bcurto-circuitos\b", re.I), "curtos-circuitos"),
    (re.compile(r"\bmeio-fios\b", re.I), "meios-fios"),
    (re.compile(r"\bbocas?-de-lobos\b", re.I), "bocas-de-lobo"),
    (re.compile(r"\bm[aã]o-francesas\b", re.I), "mãos-francesas"),
    (re.compile(r"\b[aá]gua-furtadas\b", re.I), "águas-furtadas"),
    (re.compile(r"\b(?:m[aã]o-de-obra|mao\s+de\s+obra)\b", re.I), "mão de obra"),
    (re.compile(r"\bpo[cç]o-de-visita\b", re.I), "poço de visita"),
    (re.compile(r"\bcorpo-de-prova\b", re.I), "corpo de prova"),
    (re.compile(r"\bbota\s+fora\b", re.I), "bota-fora"),
    (re.compile(r"\bmeio\s+fio\b", re.I), "meio-fio"),
    (re.compile(r"\b(?:pára-raios?|para-raio|para\s+raios?)\b", re.I), "para-raios"),
    (re.compile(r"\bcurto\s+circuito\b", re.I), "curto-circuito"),
    (re.compile(r"\bauto[-\s]+adens[aá]vel\b", re.I), "autoadensável"),
    (re.compile(r"\binfra[-\s]+estrutura\b", re.I), "infraestrutura"),
    (re.compile(r"\bà\s+(montante|jusante|prumo|seco|frio|quente|laser)\b", re.I),
     lambda m: f"a {m.group(1).lower()}"),
    (re.compile(r"\b(conforme|segundo)\s+à(s?)\b", re.I),
     lambda m: f"{m.group(1).lower()} a{m.group(2).lower()}"),
    (re.compile(r"\b(implica(?:m|rá|rão|ndo)?|acarreta(?:m|rá|rão|ram|ndo)?)\s+(?:em|na|no|nas|nos)\b", re.I),
     lambda m: m.group(1).lower()),
    (re.compile(r"\b(consiste(?:m)?)\s+de\b", re.I),
     lambda m: f"{m.group(1).lower()} em"),
    (re.compile(r"\bobedecer\s+a\s+(nbr|norma|especifica[cç][aã]o|prescri[cç][aã]o)\b", re.I),
     lambda m: f"obedecer à {m.group(1).lower()}"),
    (re.compile(r"\bproceder\s+a\s+(execu[cç][aã]o|concretagem|escava[cç][aã]o|instala[cç][aã]o|limpeza|demoli[cç][aã]o)\b", re.I),
     lambda m: f"proceder à {m.group(1).lower()}"),
    (re.compile(r"\b(situad[oa]s?|localizad[oa]s?)\s+à\s+(rua|avenida|av\.|pra[cç]a|travessa|estrada|rodovia)\b", re.I),
     lambda m: f"{m.group(1).lower()} na {m.group(2).lower()}"),
    (re.compile(r"\btratam-se\s+de\b", re.I), "trata-se de"),
    (re.compile(r"\bplanta\s+[aá]rea\b", re.I), "planta aérea"),
    (re.compile(r"\bsess[aã]o\s+(transversal|plena|mista|circular|retangular|de\s+vaz[aã]o)\b", re.I),
     lambda m: f"seção {m.group(1).lower()}"),
    (re.compile(r"\bcumprimento\s+(total|do\s+trecho|do\s+tubo|de\s+ancoragem)\b", re.I),
     lambda m: f"comprimento {m.group(1).lower()}"),
    (re.compile(r"\bhaja\s+visto\b", re.I), "haja vista"),
    (re.compile(r"\bafim\s+de\b", re.I), "a fim de"),
    (re.compile(
        r"\besta\s+(instalad[oa]s?|localizad[oa]s?|previst[oa]s?|indicad[oa]s?|representad[oa]s?|"
        r"apoiad[oa]s?|embutid[oa]s?|conectad[oa]s?|assentes?|enterrad[oa]s?|executad[oa]s?|ligad[oa]s?|"
        r"em|sob|sobre|sem)\b", re.I),
     lambda m: f"está {m.group(1).lower()}"),
]

_RE_ESPACO_ANTES_PONT = re.compile(r"([A-Za-zÀ-ÖØ-öø-ÿ\d]+)\s+([,;])")
_RE_VIRGULA_COLADA = re.compile(r"\b([A-Za-zÀ-ÖØ-öø-ÿ]{2,}),([A-Za-zÀ-ÖØ-öø-ÿ]{2,})\b")
_RE_PONT_DUPLA = re.compile(r"(?<!\.)\.\.(?!\.)|,,+|;;+")
_RE_PALAVRA_DUPLA = re.compile(r"\b([A-Za-zÀ-ÖØ-öø-ÿ]{2,})\s+\1\b", re.I)
# letra solta repetida só com um espaço ou na quebra de linha: "R U A    A N T O N I O" é nome espaçado
_RE_LETRA_DUPLA = re.compile(r"\b([eEaAoO])(?:[ \t]*\n[ \t]*| )\1\b", re.I)


_RE_EXPOENTE = re.compile(r"(?:(?<=\d)|(?<=[\s(/]))([ckm]?m)([23])(?![\dA-Za-z/])")
_RE_MODELO = re.compile(r"(?<![A-Za-z])(?:[xX]{2,}|[øØ]?[xX]\.[xX])(?![A-Za-z])|\?{2,}|\bA DEFINIR\b", re.I)
_RE_DECIMAL = re.compile(r"(?<![\d.,:/])(\d+)\.(\d+)(?![\d.,:/])")
_RE_REFERENCIA = re.compile(r"\b(?:tab\.?|tabela|item|se[cç][aã]o|cap\.?|cap[ií]tulo|anexo)\s*$", re.I)  # "Tab. 9.1"
_RE_LINHA_ACO = re.compile(r"\bR\s*\(cm\)|\bbitolas?\b", re.I)    # linhas das tabelas de dobra e de bitolas


def le_padroes(caminho: Path) -> list[str]:
    """Padrões de layer (fnmatch, sem diferença de caixa), um por linha; # comenta."""
    caminho = Path(caminho)
    if not caminho.is_file():
        return []
    linhas = (l.split("#", 1)[0].strip() for l in caminho.read_text(encoding="utf-8").splitlines())
    return [l.lower() for l in linhas if l]


def problemas_numero(texto: str, layer: str, padroes: list[str]) -> list[tuple[str, list[str]]]:
    """R-ORT-018: decimal com ponto → vírgula. Fora: milhar (ponto + 3 dígitos), texto com
    ø/Ø (armadura, padrão do detalhamento) e layer em `padroes` (topografia, base, modelo)."""
    if any(fnmatch.fnmatch(layer.lower(), p) for p in padroes):
        return []
    out = [(m.group(0), [m.group(1) + {"2": "²", "3": "³"}[m.group(2)]])     # R-ORT-019
           for m in _RE_EXPOENTE.finditer(texto)]
    if "ø" in texto or "Ø" in texto or _RE_LINHA_ACO.search(texto):
        return out
    return out + [(m.group(0), [f"{m.group(1)},{m.group(2)}"]) for m in _RE_DECIMAL.finditer(texto)
                  if len(m.group(2)) != 3 and not _RE_REFERENCIA.search(texto[:m.start()])]


def problemas_frase(texto: str) -> list[tuple[str, list[str]]]:
    """R-ORT-014 (unidades SI), R-ORT-015 (expressões/hífen/contexto), R-ORT-016 (pontuação/
    repetição) e R-ORT-020 (texto de modelo esquecido). Expoente e decimal dependem do layer:
    ficam em problemas_numero."""
    if re.fullmatch(r"[\W_]*", texto):          # só pontuação: campo vazio de carimbo
        return []
    out: list[tuple[str, list[str]]] = []
    for m in _RE_MODELO.finditer(texto):
        out.append((m.group(0), ["preencher"]))
    for m in _RE_UNIDADE.finditer(texto):
        out.append((m.group(0), [f"{m.group(1)} {_UNIDADES_SI[m.group(2)]}"]))
    for padrao, alvo in _EXPRESSOES:
        for m in padrao.finditer(texto):
            certo = alvo(m) if callable(alvo) else alvo
            out.append((m.group(0), [_caixa_de(m.group(0), certo)]))
    for m in _RE_ESPACO_ANTES_PONT.finditer(texto):
        out.append((m.group(0), [f"{m.group(1)}{m.group(2)}"]))
    for m in _RE_VIRGULA_COLADA.finditer(texto):
        out.append((m.group(0), [f"{m.group(1)}, {m.group(2)}"]))
    for m in _RE_PONT_DUPLA.finditer(texto):
        out.append((m.group(0), [m.group(0)[0]]))
    for rx in (_RE_PALAVRA_DUPLA, _RE_LETRA_DUPLA):
        for m in rx.finditer(texto):
            out.append((m.group(0), [m.group(1)]))
    if texto.count("(") > texto.count(")"):
        out.append(("(", ["fechar ')'"]))
    elif texto.count(")") > texto.count("("):
        out.append((")", ["falta '('"]))
    if texto.count("[") > texto.count("]"):
        out.append(("[", ["fechar ']'"]))
    elif texto.count("]") > texto.count("["):
        out.append(("]", ["falta '['"]))
    return out


# ------------------------------------------------------------------ textos
@dataclass
class Texto:
    espaco: str          # "Model" ou o nome do layout
    texto: str
    layer: str
    tipo: str
    handle: str          # da entidade do desenho (o INSERT, para texto de bloco)
    caixa: tuple | None  # (x0, y0, x1, y1) no espaço onde está
    altura: float


def _conteudo(e) -> str | None:
    t = e.dxftype()
    if t in ("TEXT", "ATTRIB"):
        if t == "ATTRIB" and e.dxf.get("flags", 0) & 1:     # invisível
            return None
        return e.plain_text()
    if t == "MTEXT":
        return e.plain_text()
    if t == "DIMENSION":
        v = e.dxf.get("text", "")
        return None if v.strip() in ("", "<>") else plain_mtext(v.replace("<>", " "))
    return None


def _altura(e) -> float:
    t = e.dxftype()
    if t == "MTEXT":
        return e.dxf.get("char_height", 0.0)
    if t == "DIMENSION":
        o = e.override()
        return (o.get("dimtxt", 0.0) or 0.0) * (o.get("dimscale", 1.0) or 1.0)
    return e.dxf.get("height", 0.0)


def _texto(e, espaco: str, handle: str, layer: str) -> Texto | None:
    v = (_conteudo(e) or "").strip()
    if not v:
        return None
    try:
        b = bbox.extents([e])
        caixa = (b.extmin.x, b.extmin.y, b.extmax.x, b.extmax.y) if b.has_data else None
    except Exception:
        caixa = None
    return Texto(espaco, v, layer, e.dxftype(), handle, caixa, _altura(e))


def textos(doc) -> list[Texto]:
    out = []
    for layout in [doc.modelspace()] + lay.papeis(doc):
        for e in layout:
            if e.dxftype() == "INSERT":
                for v in recursive_decompose([e]):
                    camada = v.dxf.get("layer", "0")
                    t = _texto(v, layout.name, e.dxf.handle, e.dxf.layer if camada == "0" else camada)
                    if t:
                        out.append(t)
            else:
                t = _texto(e, layout.name, e.dxf.handle, e.dxf.get("layer", "0"))
                if t:
                    out.append(t)
    return out


def arquivos(alvo: Path) -> list[Path]:
    alvo = Path(alvo)
    if alvo.is_file():
        return [alvo]
    return sorted(p for p in alvo.iterdir() if p.is_file() and p.suffix.lower() == ".dwg"
                  and not p.stem.upper().endswith("-REVISAO"))


# --------------------------------------------------------- posição na folha
class NaoAnotado(Exception):
    """O texto não aparece na folha; a mensagem diz por quê."""


def _cruza(a, b) -> bool:
    return not (a[2] < b[0] or a[0] > b[2] or a[3] < b[1] or a[1] > b[3])


def na_folha(doc, t: Texto) -> tuple[str, tuple]:
    """(layout, caixa no papel) onde o texto sai impresso."""
    if t.layer in doc.layers:
        camada = doc.layers.get(t.layer)
        if camada.is_off():
            raise NaoAnotado(f"layer {t.layer} desligada")
        if camada.is_frozen():
            raise NaoAnotado(f"layer {t.layer} congelada")
    if t.caixa is None:
        raise NaoAnotado("sem posição legível")
    if t.espaco != "Model":
        return t.espaco, t.caixa
    motivos = []
    for papel in lay.papeis(doc):
        for v in papel.query("VIEWPORT"):
            if v.dxf.id == 1:
                continue
            if not lay.de_topo(v):
                motivos.append(f"viewport {v.dxf.handle} não é vista de topo")
                continue
            tela = lay.caixa_tela(v, t.caixa)            # alvo e giro da vista aplicados
            if not _cruza(tela, lay.janela_tela(v)):
                continue
            if t.layer.upper() in {n.upper() for n in v.frozen_layers}:
                motivos.append(f"layer congelada no viewport {v.dxf.handle} ({papel.name})")
                continue
            k, c, vc = v.dxf.height / v.dxf.view_height, v.dxf.center, v.dxf.view_center_point
            x0, y0, x1, y1 = tela
            return papel.name, (c.x + (x0 - vc.x) * k, c.y + (y0 - vc.y) * k,
                                c.x + (x1 - vc.x) * k, c.y + (y1 - vc.y) * k)
    raise NaoAnotado(motivos[0] if motivos else "texto do model fora de todos os viewports")


# ------------------------------------------------------------- ocorrências
@dataclass
class Ocorrencia:
    arquivo: str
    folha: str
    layer: str
    texto: str
    palavra: str
    sugestoes: list = field(default_factory=list)
    motivo: str = ""             # vazio = anotada na folha
    caixa: tuple | None = None   # no papel
    altura: float = 0.0          # no papel


def ocorrencias(arquivo: str, doc, verif: Verificador) -> list[Ocorrencia]:
    out = []
    for t in textos(doc):
        frases = problemas_frase(t.texto) + problemas_numero(t.texto, t.layer, verif.ignorar_decimal)
        cobertas = {w.lower() for trecho, _ in frases for w in _PALAVRA.findall(trecho)}
        suspeitas = [(w, verif.sugestoes(w)) for w in palavras(t.texto)
                     if w.lower() not in cobertas and verif.suspeita(w)]
        itens = frases + suspeitas
        if not itens:
            continue
        try:
            folha, caixa = na_folha(doc, t)
            diag = lambda b: math.hypot(b[2] - b[0], b[3] - b[1])    # não muda com o giro
            altura = t.altura * (diag(caixa) / diag(t.caixa) if diag(t.caixa) else 1.0)
            motivo = ""
        except NaoAnotado as m:
            folha, caixa, altura, motivo = t.espaco, None, 0.0, str(m)
        out += [Ocorrencia(arquivo, folha, t.layer, t.texto, w, sug, motivo, caixa, altura)
                for w, sug in itens]
    return out


def grava_excel(arquivo: Path, oc: list[Ocorrencia]) -> None:
    arquivo = Path(arquivo)
    if arquivo.exists():
        raise FileExistsError(f"{arquivo.name} já existe")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Palavras"
    ws.append(["Palavra", "Ocorrências", "Sugestões", "Aceitar"])
    contagem: dict[str, list] = {}
    for o in oc:
        contagem.setdefault(o.palavra, [0, o.sugestoes])[0] += 1
    for w, (n, sug) in sorted(contagem.items(), key=lambda kv: (-kv[1][0], kv[0])):
        ws.append([w, n, ", ".join(sug), None])
    wo = wb.create_sheet("Ocorrências")
    wo.append(["Arquivo", "Folha/Model", "Layer", "Texto", "Palavra", "Sugestões", "Anotada", "Motivo"])
    for o in oc:
        wo.append([o.arquivo, o.folha, o.layer, o.texto, o.palavra, ", ".join(o.sugestoes),
                   "não" if o.motivo else "sim", o.motivo])
    for planilha, larguras in ((ws, (24, 12, 40, 10)), (wo, (32, 18, 18, 60, 22, 30, 9, 50))):
        planilha.auto_filter.ref = planilha.dimensions
        planilha.freeze_panes = "A2"
        for i, w in enumerate(larguras):
            planilha.column_dimensions[chr(65 + i)].width = w
    wb.save(arquivo)


def aprender(relatorio: Path, glossario: Path = GLOSSARIO) -> list[str]:
    ws = openpyxl.load_workbook(relatorio, read_only=True)["Palavras"]
    ja = le_glossario(glossario)
    novas = []
    for palavra, _, _, aceitar, *_ in ws.iter_rows(min_row=2, values_only=True):
        if palavra and aceitar not in (None, "") and palavra.lower() not in ja:
            ja.add(palavra.lower())
            novas.append(palavra)
    if novas:
        glossario = Path(glossario)
        glossario.parent.mkdir(parents=True, exist_ok=True)
        antes = glossario.read_text(encoding="utf-8") if glossario.is_file() else ""
        sep = "" if not antes or antes.endswith("\n") else "\n"
        glossario.write_text(antes + sep + "\n".join(novas) + "\n", encoding="utf-8")
    return novas


# --------------------------------------------------------------- revisão DWG
def destino_revisao(dwg: Path, substituir: bool) -> Path:
    destino = Path(dwg).with_name(f"{Path(dwg).stem}-REVISAO.dwg")
    if destino.exists() and not substituir:
        raise FileExistsError(f"{destino.name} já existe; use --substituir para refazer")
    return destino


def plano_anotacoes(oc: list[Ocorrencia]) -> list[dict]:
    """Uma caixa e uma nota por texto anotável (as palavras do mesmo texto vão juntas)."""
    grupos: dict[tuple, dict] = {}
    for o in oc:
        if o.motivo:
            continue
        g = grupos.setdefault((o.folha, o.caixa), dict(folha=o.folha, caixa=o.caixa, altura=o.altura, itens=[]))
        item = f"{o.palavra} → {o.sugestoes[0] if o.sugestoes else '?'}"
        if item not in g["itens"]:
            g["itens"].append(item)
    return [dict(folha=g["folha"], caixa=g["caixa"], altura=g["altura"], nota="; ".join(g["itens"]))
            for g in grupos.values()]


def _par(codigo: int, s: str) -> str:
    lit = lay.texto_lisp(s)
    return f"'({codigo} . {lit})" if lit.startswith('"') else f"(cons {codigo} {lit})"


def _mtext_ascii(s: str) -> str:
    """MTEXT aceita \\U+XXXX: assim qualquer caractere chega pelo script ASCII."""
    return "".join(ch if ord(ch) < 128 else f"\\U+{ord(ch):04X}" for ch in s)


def lisp_anotar(plano: list[dict]) -> str:
    linhas = [f"(setq camada {lay.texto_lisp(CAMADA)})",
              '(if (not (tblsearch "LAYER" camada)) (entmake (list \'(0 . "LAYER") '
              '\'(100 . "AcDbSymbolTableRecord") \'(100 . "AcDbLayerTableRecord") (cons 2 camada) '
              f'\'(70 . 0) \'(62 . {COR}) \'(6 . "Continuous"))))']
    for p in plano:
        x0, y0, x1, y1 = p["caixa"]
        h = p["altura"] or (y1 - y0) or 1.0
        m = 0.25 * h
        x0, y0, x1, y1 = x0 - m, y0 - m, x1 + m, y1 + m
        folha = _par(410, p["folha"])
        linhas.append(
            f"(entmake (list '(0 . \"LWPOLYLINE\") '(100 . \"AcDbEntity\") (cons 8 camada) {folha} "
            f"'(100 . \"AcDbPolyline\") '(90 . 4) '(70 . 1) '(10 {x0!r} {y0!r}) '(10 {x1!r} {y0!r}) "
            f"'(10 {x1!r} {y1!r}) '(10 {x0!r} {y1!r})))")
        linhas.append(
            f"(entmake (list '(0 . \"MTEXT\") '(100 . \"AcDbEntity\") (cons 8 camada) {folha} "
            f"'(100 . \"AcDbMText\") '(10 {x0!r} {y1 + m!r} 0.0) '(40 . {h!r}) '(71 . 7) "
            f"(cons 1 {lay.texto_lisp(_mtext_ascii(p['nota']))})))")
    return "\n".join(linhas) + "\n"


def confere_revisao(fonte, nova, n_entidades: int) -> list[str]:
    problemas, novas, alteradas = [], [], 0
    espacos = ["Model"] + [p.name for p in lay.papeis(fonte)]
    for nome in espacos:
        a = lay._mapa(fonte.modelspace() if nome == "Model" else fonte.layouts.get(nome))
        b = lay._mapa(nova.modelspace() if nome == "Model" else nova.layouts.get(nome))
        alteradas += sum(1 for h in a if h not in b or not lay.igual(a[h], b[h]))
        novas += [b[h] for h in set(b) - set(a)]
    if alteradas:
        problemas.append(f"{alteradas} entidade(s) existente(s) mudou/sumiu")
    fora = [e for e in novas if e.dxf.layer != CAMADA]
    if fora:
        problemas.append(f"{len(fora)} entidade(s) nova(s) fora da layer {CAMADA}")
    if len(novas) != n_entidades:
        problemas.append(f"{len(novas)} anotação(ões) no desenho, planejadas {n_entidades}")
    return problemas


def revisar(alvo: Path, oda, core, glossario: Path = GLOSSARIO, substituir: bool = False,
            *, relatorio: Path | None = None, armadilhas: Path = ARMADILHAS,
            ignorar_decimal: Path = IGNORAR_DECIMAL) -> dict:
    alvo = Path(alvo)
    arqs = arquivos(alvo)
    if not arqs:
        raise ValueError(f"nenhum DWG em {alvo}")
    pasta = alvo if alvo.is_dir() else alvo.parent
    relatorio = Path(relatorio) if relatorio is not None else pasta / f"ORTOGRAFIA {datetime.date.today():%Y-%m-%d}.xlsx"
    if relatorio.exists():
        raise FileExistsError(f"{relatorio.name} já existe")
    verif = Verificador(le_glossario(glossario), armadilhas=le_armadilhas(armadilhas),
                        ignorar_decimal=le_padroes(ignorar_decimal))
    trabalho = Path(tempfile.mkdtemp(prefix="toolkit-ortografia-"))
    todas, planos = [], {}
    for a in arqs:
        doc = ezdxf.readfile(lay._le(a, oda, trabalho))
        oc = ocorrencias(a.name, doc, verif)
        todas += oc
        plano = plano_anotacoes(oc)
        if plano:
            planos[a] = (doc, plano)
    destinos = {a: destino_revisao(a, substituir) for a in planos}    # para antes de gravar
    grava_excel(relatorio, todas)
    revisoes = []
    for i, (a, (doc, plano)) in enumerate(planos.items()):
        tmp = trabalho / f"revisao-{i}.dwg"
        lay.roda_core(core, a, lisp_anotar(plano), tmp, trabalho)
        problemas = confere_revisao(doc, ezdxf.readfile(lay.dwg_para_dxf(oda, tmp, trabalho)), 2 * len(plano))
        if problemas:
            raise RuntimeError(f"{destinos[a].name} não conferiu, não entreguei: " + "; ".join(problemas)
                               + f". Consulte {trabalho}.")
        shutil.copy2(tmp, destinos[a])
        revisoes.append(destinos[a])
    return dict(relatorio=relatorio, revisoes=revisoes, ocorrencias=todas)


# -------------------------------------------------------------------- CLI
def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--glossario", type=Path, default=GLOSSARIO, help=f"padrão: {GLOSSARIO}")
    p.add_argument("--armadilhas", type=Path, default=ARMADILHAS, help=f"padrão: {ARMADILHAS}")
    p.add_argument("--ignorar-decimal", type=Path, default=IGNORAR_DECIMAL, help=f"padrão: {IGNORAR_DECIMAL}")
    sub = p.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("revisar", help="Relatório Excel + cópias <nome>-REVISAO.dwg anotadas")
    r.add_argument("alvo", type=Path, help="Pasta (DWG da raiz) ou um DWG")
    r.add_argument("--substituir", action="store_true", help="Refaz cópias -REVISAO existentes")
    r.add_argument("--core", help="Caminho do accoreconsole.exe")
    r.add_argument("--oda", help="Caminho do ODAFileConverter.exe")
    a_ = sub.add_parser("aprender", help="Passa ao glossário as palavras com Aceitar marcado")
    a_.add_argument("relatorio", type=Path)
    a = p.parse_args(argv)
    try:
        if a.cmd == "aprender":
            novas = aprender(a.relatorio, a.glossario)
            print(f"{len(novas)} palavra(s) no glossário {a.glossario}: {', '.join(novas)}")
            return 0
        oda = lay.localizar(a.oda, "ODA/ODAFileConverter */ODAFileConverter.exe", "ODA File Converter")
        core = lay.localizar(a.core, "Autodesk/AutoCAD 20[0-9][0-9]/accoreconsole.exe", "accoreconsole")
        res = revisar(a.alvo, oda, core, a.glossario, a.substituir, armadilhas=a.armadilhas,
                      ignorar_decimal=a.ignorar_decimal)
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError) as erro:
        print(f"ERRO: {erro}", file=sys.stderr)
        return 1
    oc = res["ocorrencias"]
    print(f"{len({o.palavra for o in oc})} palavra(s) suspeita(s) em {len(oc)} ocorrência(s); "
          f"{sum(1 for o in oc if not o.motivo)} anotada(s) na folha")
    print(f"relatório: {res['relatorio']}")
    for d in res["revisoes"]:
        print(f"revisão:   {d}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

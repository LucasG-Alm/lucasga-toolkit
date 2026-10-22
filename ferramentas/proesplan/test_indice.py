"""Índice de desenhos (Word) a partir dos carimbos dos DWG da pasta.

Carimbo e Word sintéticos; integração (-m autocad) lê DWG de verdade pelo ODA.
"""
import datetime

import docx
import ezdxf
import pytest
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm

import indice as ind
import layouts as lay

TAGS = [  # (tag, y): a altura no carimbo define a ordem das linhas do título
    ("TÍTULO1", 47.0), ("TÍTULO2", 45.9), ("TÍTULO1-UMA_LINHA", 43.0), ("TÍTULO3", 41.5),
    ("TÍTULO4", 34.7), ("REVISÃO", 31.7), ("FOLHA", 31.7), ("TÍTULO5", 30.4), ("TÍTULO6", 25.2),
    ("NÚM.CONTRATADA", 21.2), ("TÍTULO7", 20.9), ("DESENHISTA", 10), ("PROJETISTA", 9),
    ("APROVADO_POR", 8), ("ÁREA-PROJETO", 7), ("OCULTO", 6),
]


def _carimbo_doc(folhas: dict):
    """{layout: {tag: valor}} — um carimbo por layout."""
    doc = ezdxf.new("R2018")
    b = doc.blocks.new("145-PAD-A1")
    for tag, y in TAGS:
        b.add_attdef(tag, (700, y), dxfattribs={"height": 2.5, "flags": 1 if tag == "OCULTO" else 0})
    for nome, valores in folhas.items():
        papel = doc.layouts.new(nome)
        papel.reset_main_viewport()
        i = papel.add_blockref("145-PAD-A1", (0, 0))
        i.add_auto_attribs({tag: valores.get(tag, "") for tag, _ in TAGS})
    doc.layouts.delete("Layout1")
    return doc


TOP = {"NÚM.CONTRATADA": "900-TOP-RES-001", "REVISÃO": "00", "FOLHA": "01/01",
       "TÍTULO2": "PROJETO EXECUTIVO DO CENTRO DE RESERVAÇÃO DO",
       "TÍTULO3": "EMPREENDIMENTO FICTÍCIO DE EXEMPLO", "TÍTULO4": "LEVANTAMENTO PLANIALTIMÉTRICO",
       "TÍTULO5": ".", "TÍTULO6": "TÍTULO6", "TÍTULO7": "PLANTA GERAL", "OCULTO": "NÃO APARECE"}


# ----------------------------------------------------------------- carimbo
def test_numero_revisao_folha_e_titulo_de_cima_para_baixo():
    """R-IND-002"""
    folhas, avisos = ind.folhas(_carimbo_doc({"900-TOP-RES-001": TOP}), "top.dwg")
    assert avisos == []
    f = folhas[0]
    assert (f.numero, f.revisao, f.folha) == ("900-TOP-RES-001", "Rev-0", "01/01")
    assert f.titulo == ["PROJETO EXECUTIVO DO CENTRO DE RESERVAÇÃO DO",
                        "EMPREENDIMENTO FICTÍCIO DE EXEMPLO", "LEVANTAMENTO PLANIALTIMÉTRICO",
                        "PLANTA GERAL"]                  # sem '.', placeholder nem invisível
    assert (f.arquivo, f.layout) == ("top.dwg", "900-TOP-RES-001")


def test_revisao_nao_numerica_fica_como_esta():
    """R-IND-002"""
    folhas, _ = ind.folhas(_carimbo_doc({"X": dict(TOP, **{"REVISÃO": "A"})}), "x.dwg")
    assert folhas[0].revisao == "Rev-A"


def test_layout_sem_carimbo_e_avisado():
    """R-IND-001"""
    doc = _carimbo_doc({"900-TOP-RES-001": TOP})
    vazio = doc.layouts.new("RASCUNHO")
    vazio.add_line((0, 0), (1, 1))
    folhas, avisos = ind.folhas(doc, "top.dwg")
    assert len(folhas) == 1 and avisos and "RASCUNHO" in avisos[0]


def test_dwg_da_raiz_sem_revisoes(tmp_path):
    """R-IND-001"""
    for n in ("A.dwg", "A-REVISAO.dwg", "x.docx"):
        (tmp_path / n).write_bytes(b"")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "B.dwg").write_bytes(b"")
    assert [p.name for p in ind.desenhos(tmp_path)] == ["A.dwg"]


# ------------------------------------------------------------------- ordem
def _f(numero):
    return ind.Folha(numero, "Rev-0", "01/01", ["T"], "a.dwg", numero)


def test_ordem_das_disciplinas_segue_o_indice_e_nova_vai_ao_fim():
    """R-IND-003"""
    fs = [_f(n) for n in ("900-EST-RES-002", "900-HID-RES-001", "900-ELE-RES-001",
                          "900-TOP-RES-001", "900-EST-RES-001", "900-AAA-RES-001")]
    ordem = ["TOP", "GER", "HID", "EST"]
    assert [f.numero for f in ind.ordena(fs, ordem)] == [
        "900-TOP-RES-001", "900-HID-RES-001", "900-EST-RES-001", "900-EST-RES-002",
        "900-AAA-RES-001", "900-ELE-RES-001"]


def test_numero_repetido_ou_vazio_para():
    """R-IND-004"""
    with pytest.raises(ValueError, match="repetido"):
        ind.valida([_f("900-TOP-RES-001"), _f("900-TOP-RES-001")])
    with pytest.raises(ValueError, match="NÚM.CONTRATADA"):
        ind.valida([_f("")])


# -------------------------------------------------------------------- Word
def _indice(caminho):
    d = docx.Document()
    d.add_paragraph("RELAÇÃO DE DESENHOS")
    t = d.add_table(rows=3, cols=3)
    for c, txt in zip(t.rows[0].cells, ("NÚMERO", "TÍTULO", "FL.")):
        c.text = txt
    linha = t.rows[1]
    linha.height, linha.height_rule = Cm(2.5), WD_ROW_HEIGHT_RULE.EXACTLY
    t.rows[2].height, t.rows[2].height_rule = Cm(0.5), WD_ROW_HEIGHT_RULE.EXACTLY

    def para(celula, texto, primeiro=False, **fmt):
        p = celula.paragraphs[0] if primeiro else celula.add_paragraph()
        r = p.add_run(texto)
        r.font.name = "Arial"
        r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
        for k, v in fmt.items():
            setattr(r.font, k, v)
    para(linha.cells[0], "900-TOP-RES-001", True)
    para(linha.cells[0], "Rev-0", italic=True)
    para(linha.cells[1], "LINHA 1", True, small_caps=True)
    para(linha.cells[1], "LINHA 2", small_caps=True)
    para(linha.cells[1], "LINHA 3")
    para(linha.cells[1], "LINHA 4")                 # o índice real tem 4 linhas de título
    para(linha.cells[2], "01/01", True)
    d.add_paragraph("depois da tabela")
    d.save(caminho)
    return caminho


FOLHAS = [ind.Folha("900-TOP-RES-001", "Rev-0", "01/01", ["A", "B", "C", "D"], "t.dwg", "t"),
          ind.Folha("900-EST-RES-004", "Rev-1", "01/02", ["A", "B", "C", "D", "E", "F"], "e.dwg", "e")]


def _textos(celula):
    return [p.text for p in celula.paragraphs]


def test_disciplinas_lidas_do_indice_existente(tmp_path):
    """R-IND-003"""
    assert ind.disciplinas(docx.Document(_indice(tmp_path / "i.docx"))) == ["TOP"]


def test_tabela_refeita_com_cabecalho_espacador_e_o_resto_intacto(tmp_path):
    """R-IND-005"""
    d = docx.Document(_indice(tmp_path / "i.docx"))
    ind.preenche(d, FOLHAS)
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["NÚMERO", "TÍTULO", "FL."]
    assert len(t.rows) == 1 + 2 * len(FOLHAS)
    assert _textos(t.rows[1].cells[0]) == ["900-TOP-RES-001", "Rev-0"]
    assert _textos(t.rows[1].cells[1]) == ["A", "B", "C", "D"]
    assert _textos(t.rows[3].cells[2]) == ["01/02"]
    assert all(c.text == "" for c in t.rows[2].cells)                  # espaçador
    assert [p.text for p in d.paragraphs] == ["RELAÇÃO DE DESENHOS", "depois da tabela"]


def test_formatacao_clonada_do_modelo_sem_marca_texto(tmp_path):
    """R-IND-005"""
    d = docx.Document(_indice(tmp_path / "i.docx"))
    ind.preenche(d, FOLHAS)
    titulo = d.tables[0].rows[3].cells[1].paragraphs
    assert [bool(p.runs[0].font.small_caps) for p in titulo] == [True, True, False, False, False, False]
    assert d.tables[0].rows[1].cells[0].paragraphs[1].runs[0].font.italic
    assert all(r.font.name == "Arial" for p in titulo for r in p.runs)
    assert not any(r.font.highlight_color for row in d.tables[0].rows for c in row.cells
                   for p in c.paragraphs for r in p.runs)


def test_titulo_maior_que_o_modelo_vira_altura_minima(tmp_path):
    """R-IND-005"""
    d = docx.Document(_indice(tmp_path / "i.docx"))
    ind.preenche(d, FOLHAS)
    rows = d.tables[0].rows
    assert rows[1].height_rule == WD_ROW_HEIGHT_RULE.EXACTLY       # 4 linhas: cabe
    assert rows[3].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST      # 6 linhas: não cabe


# ------------------------------------------------------------ gravação
def test_backup_antes_de_sobrescrever(tmp_path):
    """R-IND-006"""
    arq = _indice(tmp_path / "900-INDICE.docx")
    antes = arq.read_bytes()
    agora = datetime.datetime(2026, 10, 21, 18, 5)
    backup = ind.grava(docx.Document(arq), arq, agora)
    assert backup == tmp_path / "_BACKUP" / "900-INDICE (antes de atualizar 2026-10-21 1805).docx"
    assert backup.read_bytes() == antes


def test_word_aberto_para(tmp_path):
    """R-IND-006"""
    arq = _indice(tmp_path / "900-INDICE.docx")
    (tmp_path / "~$0-INDICE.docx").write_bytes(b"")
    with pytest.raises(OSError, match="aberto"):
        ind.grava(docx.Document(arq), arq)


def test_acha_o_indice_da_pasta(tmp_path):
    """R-IND-006"""
    _indice(tmp_path / "900-INDICE DE DESENHOS-R00.docx")
    assert ind.acha_indice(tmp_path).name == "900-INDICE DE DESENHOS-R00.docx"
    _indice(tmp_path / "OUTRO INDICE.docx")
    with pytest.raises(ValueError, match="mais de um"):
        ind.acha_indice(tmp_path)


# ------------------------------------------------------ integração (ODA)
@pytest.mark.autocad
def test_integracao_atualiza_o_indice_pelos_dwg(tmp_path):
    """R-IND-001, R-IND-002, R-IND-005, R-IND-006"""
    try:
        oda = lay.localizar(None, "ODA/ODAFileConverter */ODAFileConverter.exe", "ODA")
        core = lay.localizar(None, "Autodesk/AutoCAD 20[0-9][0-9]/accoreconsole.exe", "AutoCAD")
    except FileNotFoundError:
        pytest.skip("ODA ou AutoCAD Core Console não instalados")
    obra = tmp_path / "Desenhos"
    obra.mkdir()
    est = dict(TOP, **{"NÚM.CONTRATADA": "900-EST-RES-001", "TÍTULO7": "FORMAS"})
    for nome, folhas in (("top", {"900-TOP-RES-001": TOP}), ("est", {"900-EST-RES-001": est})):
        _carimbo_doc(folhas).saveas(tmp_path / f"{nome}.dxf")
        lay.roda_core(core, tmp_path / f"{nome}.dxf", "", obra / f"{nome}.dwg", tmp_path)
    _indice(obra / "900-INDICE DE DESENHOS-R00.docx")
    r = ind.atualizar(obra, oda)
    assert r["backup"].is_file()
    t = docx.Document(r["indice"]).tables[0]
    assert [t.rows[i].cells[0].paragraphs[0].text for i in (1, 3)] == ["900-TOP-RES-001", "900-EST-RES-001"]
    assert t.rows[3].cells[1].paragraphs[-1].text == "FORMAS"

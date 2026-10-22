"""Ortografia dos textos dos desenhos: dicionário VERO, glossário, Excel e DWG de revisão.

Lógica e leitura com DXF sintético; integração (-m autocad) com Core Console e ODA.
"""
import ezdxf
import openpyxl
import pytest

import layouts as lay
import ortografia as ort


@pytest.fixture(scope="module")
def verif():
    return ort.Verificador(glossario={"proesplan", "ESC"})


# --------------------------------------------------------------- palavras
def test_palavras_ignora_trecho_com_digito_e_letra_solta():
    """R-ORT-003"""
    assert ort.palavras("RESERVATÓRIO 8,00X8,00M - AÇO CA-50 (ESC. 1:50) e a") == \
        ["RESERVATÓRIO", "AÇO", "ESC"]


def test_palavras_separa_hifen_e_barra():
    """R-ORT-003"""
    assert ort.palavras("PLANTAS/CORTES guarda-corpo") == ["PLANTAS", "CORTES", "guarda", "corpo"]


# ------------------------------------------------------------ dicionário
def test_maiuscula_e_acento_valem_como_no_dicionario(verif):
    """R-ORT-004"""
    assert not verif.suspeita("FUNDAÇÃO")
    assert not verif.suspeita("Detalhamento")
    assert verif.suspeita("FUNDACAO")


def test_glossario_aceita_sem_diferenca_de_caixa(verif):
    """R-ORT-005"""
    assert not verif.suspeita("PROESPLAN")
    assert not verif.suspeita("esc")


def test_glossario_txt_uma_palavra_por_linha_com_comentario(tmp_path):
    """R-ORT-005"""
    g = tmp_path / "g.txt"
    g.write_text("# técnicos\nGrauteamento\n\n  Fck  \n", encoding="utf-8")
    assert ort.le_glossario(g) == {"grauteamento", "fck"}
    assert ort.le_glossario(tmp_path / "nao_existe.txt") == set()


def test_sugestao_restaura_acento_na_caixa_original(verif):
    """R-ORT-006"""
    assert verif.sugestoes("FUNDACAO") == ["FUNDAÇÃO"]
    assert verif.sugestoes("reservatorio") == ["reservatório"]


def test_sugestao_a_uma_letra_quando_nao_e_acento(verif):
    """R-ORT-006"""
    assert "DETALHAMENTO" in verif.sugestoes("DETALHAMNETO")


def test_sem_candidata_fica_sem_sugestao(verif):
    """R-ORT-006"""
    assert verif.sugestoes("XQWZKJ") == []


# ---------------------------------------------------------------- textos
def _doc():
    doc = ezdxf.new("R2018")
    doc.layers.add("TXT")
    doc.layers.add("CAR")
    m = doc.modelspace()
    m.add_text("FUNDACAO DO MODEL", height=2, dxfattribs={"layer": "TXT"}).set_placement((10, 10))
    m.add_text("LONGE", height=2, dxfattribs={"layer": "TXT"}).set_placement((900, 900))
    m.add_mtext("{\\fArial|b1;CONCRETO}\\PARMADURA", dxfattribs={"layer": "TXT", "insert": (10, 40),
                                                                   "char_height": 2})
    b = doc.blocks.new("CARIMBO")
    b.add_text("DESENHSTA FIXO", height=2).set_placement((0, 0))
    b.add_attdef("TITULO", (0, 5), dxfattribs={"height": 2})
    b.add_attdef("OCULTO", (0, 10), dxfattribs={"height": 2, "flags": 1})
    papel = doc.layouts.new("F-01")
    papel.reset_main_viewport()
    papel.add_viewport(center=(150, 100), size=(100, 100), view_center_point=(50, 50), view_height=100)
    i = papel.add_blockref("CARIMBO", (200, 20), dxfattribs={"layer": "CAR"})
    i.add_auto_attribs({"TITULO": "PLANTA GERAL", "OCULTO": "ESCONDIDO"})
    doc.layouts.delete("Layout1")
    return doc


def test_textos_de_papel_model_blocos_e_atributos_visiveis():
    """R-ORT-002"""
    t = {(x.espaco, x.texto) for x in ort.textos(_doc())}
    assert ("Model", "FUNDACAO DO MODEL") in t
    assert ("Model", "CONCRETO\nARMADURA") in t                  # MTEXT sem formatação
    assert ("F-01", "DESENHSTA FIXO") in t and ("F-01", "PLANTA GERAL") in t
    assert not any("ESCONDIDO" in x for _, x in t)                # atributo invisível


def test_texto_de_cota_so_quando_digitado():
    """R-ORT-002"""
    doc = ezdxf.new("R2018", setup=True)
    m = doc.modelspace()
    m.add_linear_dim(base=(0, 5), p1=(0, 0), p2=(10, 0)).render()
    m.add_linear_dim(base=(0, 15), p1=(0, 10), p2=(10, 10), text="VAO LIVRE").render()
    textos = [x.texto for x in ort.textos(doc) if x.tipo == "DIMENSION"]
    assert textos == ["VAO LIVRE"]


def test_arquivos_da_raiz_sem_as_revisoes_nem_subpastas(tmp_path):
    """R-ORT-001"""
    for n in ("A.dwg", "B.DWG", "A-REVISAO.dwg", "nota.txt"):
        (tmp_path / n).write_bytes(b"")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "C.dwg").write_bytes(b"")
    assert [p.name for p in ort.arquivos(tmp_path)] == ["A.dwg", "B.DWG"]
    assert ort.arquivos(tmp_path / "A.dwg") == [tmp_path / "A.dwg"]


# ------------------------------------------------------- posição na folha
def _texto(doc, conteudo):
    return next(x for x in ort.textos(doc) if x.texto == conteudo)


def test_texto_do_papel_fica_onde_esta():
    """R-ORT-011"""
    doc = _doc()
    t = _texto(doc, "PLANTA GERAL")
    assert ort.na_folha(doc, t) == ("F-01", t.caixa)


def test_texto_do_model_vai_para_o_papel_pelo_viewport():
    """R-ORT-011: viewport 100x100 no papel mostrando 100x100 do model, centro (50,50)→(150,100)."""
    doc = _doc()
    t = _texto(doc, "FUNDACAO DO MODEL")
    folha, (x0, y0, x1, y1) = ort.na_folha(doc, t)
    assert folha == "F-01"
    assert (x0, y0) == pytest.approx((t.caixa[0] + 100, t.caixa[1] + 50))
    assert x1 - x0 == pytest.approx(t.caixa[2] - t.caixa[0])


def test_texto_do_model_por_viewport_girado_e_com_alvo():
    """R-ORT-011. Regressão: 'PONTOS CRITIOS' saía no PDF do 900-GER-LOC-001 (viewport com
    giro 270° e alvo deslocado) e a ferramenta dizia que o texto estava fora da folha."""
    doc = _doc()
    v = next(iter(doc.layouts.get("F-01").query("VIEWPORT[id>1]")))
    v.dxf.view_target_point, v.dxf.view_twist_angle = (0, 0, 0), 90.0
    v.dxf.view_center_point = (-10, 10)              # giro 90° de (10, 10): a inserção do texto
    folha, (x0, y0, x1, y1) = ort.na_folha(doc, _texto(doc, "FUNDACAO DO MODEL"))
    assert folha == "F-01"
    assert x0 - 1e-6 <= 150 <= x1 + 1e-6 and y0 - 1e-6 <= 100 <= y1 + 1e-6   # inserção → centro do viewport
    assert (y1 - y0) > (x1 - x0)                      # texto deitado no model fica em pé na folha


@pytest.mark.parametrize("mexe,motivo", [
    (lambda d: None, "fora"),
    (lambda d: d.layers.get("TXT").off(), "desligada"),
    (lambda d: d.layers.get("TXT").freeze(), "congelada"),
])
def test_texto_que_nao_sai_na_folha_traz_o_motivo(mexe, motivo):
    """R-ORT-011"""
    doc = _doc()
    mexe(doc)
    conteudo = "LONGE" if motivo == "fora" else "FUNDACAO DO MODEL"
    with pytest.raises(ort.NaoAnotado, match=motivo):
        ort.na_folha(doc, _texto(doc, conteudo))


# ------------------------------------------------------------------ Excel
def _ocorrencias(verif):
    doc = _doc()
    return ort.ocorrencias("obra.dwg", doc, verif)


def test_ocorrencias_trazem_palavra_sugestao_e_motivo(verif):
    """R-ORT-007, R-ORT-011"""
    oc = _ocorrencias(verif)
    fund = next(o for o in oc if o.palavra == "FUNDACAO")
    assert fund.sugestoes == ["FUNDAÇÃO"] and fund.folha == "F-01" and fund.motivo == ""
    assert {o.palavra for o in oc} >= {"FUNDACAO", "DESENHSTA"}


def test_excel_tem_palavras_com_aceitar_e_ocorrencias(tmp_path, verif):
    """R-ORT-007"""
    arq = tmp_path / "ORTOGRAFIA.xlsx"
    ort.grava_excel(arq, _ocorrencias(verif))
    wb = openpyxl.load_workbook(arq)
    palavras = [list(r) for r in wb["Palavras"].iter_rows(values_only=True)]
    assert palavras[0][:4] == ["Palavra", "Ocorrências", "Sugestões", "Aceitar"]
    assert ["FUNDACAO", 1, "FUNDAÇÃO", None] in [p[:4] for p in palavras]
    cab = [c.value for c in wb["Ocorrências"][1]]
    assert cab[:5] == ["Arquivo", "Folha/Model", "Layer", "Texto", "Palavra"]
    with pytest.raises(FileExistsError):
        ort.grava_excel(arq, [])


def test_aprender_acrescenta_marcadas_sem_repetir(tmp_path, verif):
    """R-ORT-008"""
    arq = tmp_path / "r.xlsx"
    ort.grava_excel(arq, _ocorrencias(verif))
    wb = openpyxl.load_workbook(arq)
    ws = wb["Palavras"]
    for linha in ws.iter_rows(min_row=2):
        if linha[0].value in ("DESENHSTA", "FUNDACAO"):
            linha[3].value = "x"
    wb.save(arq)
    g = tmp_path / "glossario.txt"
    g.write_text("desenhsta\n", encoding="utf-8")
    assert ort.aprender(arq, g) == ["FUNDACAO"]
    assert ort.le_glossario(g) == {"desenhsta", "fundacao"}


# ------------------------------------------------------------- revisão DWG
def test_copia_de_revisao_existente_para_sem_substituir(tmp_path):
    """R-ORT-009"""
    (tmp_path / "A-REVISAO.dwg").write_bytes(b"")
    with pytest.raises(FileExistsError, match="A-REVISAO"):
        ort.destino_revisao(tmp_path / "A.dwg", substituir=False)
    assert ort.destino_revisao(tmp_path / "A.dwg", substituir=True) == tmp_path / "A-REVISAO.dwg"


def test_plano_uma_caixa_e_uma_nota_por_texto(verif):
    """R-ORT-010"""
    plano = ort.plano_anotacoes(_ocorrencias(verif))
    fund = next(p for p in plano if "FUNDACAO" in p["nota"])
    assert fund["folha"] == "F-01" and "FUNDAÇÃO" in fund["nota"] and fund["altura"] == pytest.approx(2)
    assert len({(p["folha"], p["caixa"]) for p in plano}) == len(plano)    # um por texto


def test_lisp_cria_layer_210_e_desenha_no_papel_da_folha(verif):
    """R-ORT-010"""
    lsp = ort.lisp_anotar(ort.plano_anotacoes(_ocorrencias(verif)))
    assert "(62 . 210)" in lsp and '(410 . "F-01")' in lsp
    assert "LWPOLYLINE" in lsp and "MTEXT" in lsp
    assert "(chr 199)" in lsp                       # Ç de _ANOTAÇÕES montado em ASCII
    lsp.encode("ascii")


def test_conferencia_aceita_so_anotacoes_novas():
    """R-ORT-012"""
    fonte = _doc()
    nova = _doc()
    nova.layers.add(ort.CAMADA, color=210)
    nova.layouts.get("F-01").add_lwpolyline([(0, 0), (1, 1)], dxfattribs={"layer": ort.CAMADA})
    plano = [dict(folha="F-01", caixa=(0, 0, 1, 1), nota="x", altura=1)]
    assert ort.confere_revisao(fonte, nova, n_entidades=1) == []
    nova.layouts.get("F-01").add_line((0, 0), (5, 5), dxfattribs={"layer": "TXT"})
    assert ort.confere_revisao(fonte, nova, n_entidades=1)
    alterada = _doc()
    next(iter(alterada.modelspace().query("TEXT"))).dxf.text = "OUTRO"
    assert ort.confere_revisao(fonte, alterada, n_entidades=0)


# ------------------------------------------------------ integração (AutoCAD)
@pytest.mark.autocad
def test_integracao_relatorio_e_copia_anotada(tmp_path):
    """R-ORT-001, R-ORT-007, R-ORT-009, R-ORT-010, R-ORT-012"""
    try:
        oda = lay.localizar(None, "ODA/ODAFileConverter */ODAFileConverter.exe", "ODA")
        core = lay.localizar(None, "Autodesk/AutoCAD 20[0-9][0-9]/accoreconsole.exe", "AutoCAD")
    except FileNotFoundError:
        pytest.skip("ODA ou AutoCAD Core Console não instalados")
    dxf = tmp_path / "base.dxf"
    _doc().saveas(dxf)
    obra = tmp_path / "obra"
    obra.mkdir()
    lay.roda_core(core, dxf, "", obra / "F-01.dwg", tmp_path)          # DWG de verdade
    g = tmp_path / "glossario.txt"
    g.write_text("DESENHSTA\n", encoding="utf-8")
    r = ort.revisar(obra, oda, core, glossario=g)
    assert r["relatorio"].parent == obra and r["relatorio"].suffix == ".xlsx"
    assert r["revisoes"] == [obra / "F-01-REVISAO.dwg"]
    doc = ezdxf.readfile(lay.dwg_para_dxf(oda, r["revisoes"][0], tmp_path))
    novas = [e for e in doc.layouts.get("F-01") if e.dxf.layer == ort.CAMADA]
    assert sorted(e.dxftype() for e in novas) == ["LWPOLYLINE", "MTEXT"]    # só FUNDACAO
    assert doc.layers.get(ort.CAMADA).color == 210


# ------------------------------------------------------------- armadilhas
def test_armadilha_marca_palavra_valida_sem_acento():
    """R-ORT-013: CRITICOS existe no dicionário, mas em projeto é CRÍTICOS."""
    v = ort.Verificador(armadilhas={"criticos": "críticos", "forma": "fôrma"})
    assert v.suspeita("CRITICOS") and v.sugestoes("CRITICOS") == ["CRÍTICOS"]
    assert v.suspeita("Forma") and v.sugestoes("Forma") == ["Fôrma"]
    assert not v.suspeita("PONTOS")


def test_glossario_vence_a_armadilha():
    """R-ORT-013"""
    v = ort.Verificador(glossario={"forma"}, armadilhas={"forma": "fôrma"})
    assert not v.suspeita("FORMA")


def test_arquivo_de_armadilhas_errado_igual_certo(tmp_path):
    """R-ORT-013"""
    a = tmp_path / "a.txt"
    a.write_text("# estruturas\nforma = fôrma   # concreto\n\nTITULO=título\n", encoding="utf-8")
    assert ort.le_armadilhas(a) == {"forma": "fôrma", "titulo": "título"}


def test_lista_padrao_tem_os_casos_achados_em_regressoes():
    """R-ORT-013"""
    padrao = ort.le_armadilhas(ort.ARMADILHAS)
    assert {"criticos", "titulo", "forma", "formas", "area", "agua"} <= set(padrao)
    assert "esta" not in padrao and "e" not in padrao      # dependem de contexto


def test_armadilhas_incluem_novos_acentos_e_jargoes_de_canteiro():
    """R-ORT-013, R-ORT-017"""
    padrao = ort.le_armadilhas(ort.ARMADILHAS)
    # acentos técnicos comuns
    assert padrao["nivel"] == "nível"
    assert padrao["vao"] == "vão"
    assert padrao["angulo"] == "ângulo"
    assert padrao["reservatorio"] == "reservatório"
    assert padrao["sanitario"] == "sanitário"
    assert padrao["piezometrica"] == "piezométrica"
    # desvios fonéticos e jargões de canteiro (R-ORT-017)
    assert padrao["mucheta"] == "mocheta"
    assert padrao["boeiro"] == "bueiro"
    assert padrao["estanqueiedade"] == "estanqueidade"
    assert padrao["impermeabilisação"] == "impermeabilização"
    assert padrao["imbutido"] == "embutido"
    assert padrao["cumieira"] == "cumeeira"
    assert padrao["xumbador"] == "chumbador"
    assert padrao["galbião"] == "gabião"
    assert padrao["conduite"] == "conduíte"
    # exemplos de regressão e vocabulário técnico
    assert padrao["terreo"] == "térreo"
    assert padrao["condominio"] == "condomínio"
    assert padrao["fômas"] == "fôrmas"
    assert padrao["critios"] == "críticos"
    assert padrao["detalhametno"] == "detalhamento"
    assert padrao["resistro"] == "registro"
    assert padrao["flage"] == "flange"
    assert padrao["cotuvelo"] == "cotovelo"
    assert padrao["terraplanagem"] == "terraplenagem"
    assert padrao["bascultante"] == "basculante"
    assert padrao["buxa"] == "bucha"


# ---------------------------------------------- unidades, expressões e pontuação
@pytest.mark.parametrize("trecho,esperado,sugestao", [
    ("EXTENSÃO DE 50Km NA VIA", "50Km", "50 km"),
    ("LARGURA 20 mts LIVRE", "20 mts", "20 m"),
    ("FCK = 30mpa MÍNIMO", "30mpa", "30 MPa"),
    ("PRESSÃO 150 Kpa", "150 Kpa", "150 kPa"),
    ("CARGA DE 25kn POR APOIO", "25kn", "25 kN"),
    ("VAZÃO 12,5 l/seg PROJETADA", "12,5 l/seg", "12,5 L/s"),
    ("VAZÃO 80 m3/hr", "80 m3/hr", "80 m³/h"),
    ("TEMPO 24hrs DE ENSAIO", "24hrs", "24 h"),
    ("POTÊNCIA 75kva TRIFÁSICO", "75kva", "75 kVA"),
    ("REDE DE 13,8KV AÉREA", "13,8KV", "13,8 kV"),
    ("BANCO DE 50kVAr AUTOMÁTICO", "50kVAr", "50 kvar"),
    ("CONSUMO DE 120kW/h MENSAL", "120kW/h", "120 kWh"),
    ("MASSA DE 65kgs DO MARTELO", "65kgs", "65 kg"),
    ("FREQUÊNCIA DE 60hz NOMINAL", "60hz", "60 Hz"),
])
def test_unidades_fora_do_si_sao_marcadas_com_sugestao(trecho, esperado, sugestao):
    """R-ORT-014"""
    achados = dict(ort.problemas_frase(trecho))
    assert achados.get(esperado) == [sugestao]


def test_unidades_validas_e_cotas_de_desenho_nao_disparam_falso_positivo():
    """R-ORT-014: cotas coladas válidas (8,00X8,00M, 30MPa, 110mm) ficam aceitas."""
    texto = "RESERVATÓRIO 8,00X8,00M - FCK 30 MPa - TUBO DN110mm - EXTENSÃO 50 km (ESC. 1:50)"
    assert ort.problemas_frase(texto) == []


@pytest.mark.parametrize("frase,termo,sugestao", [
    ("REGULARIZAÇÃO DO SUB-LEITO COMPACTADO", "SUB-LEITO", "SUBLEITO"),
    ("CAMADA DE SUB BASE GRANULAR", "SUB BASE", "SUB-BASE"),
    ("CAMADA DE subbase granular", "subbase", "sub-base"),
    ("EXECUÇÃO DE CONTRA-PISO ARMADO", "CONTRA-PISO", "CONTRAPISO"),
    ("BLOCO PRE-MOLDADO DE CONCRETO", "PRE-MOLDADO", "PRÉ-MOLDADO"),
    ("LAJE PRÉ MOLDADA", "PRÉ MOLDADA", "PRÉ-MOLDADA"),
    ("DETALHE DO GUARDA CORPO METÁLICO", "GUARDA CORPO", "GUARDA-CORPO"),
    ("ALTURA DE PÉ DIREITO DUPLO", "PÉ DIREITO", "PÉ-DIREITO"),
    ("CUSTO DE MÃO-DE-OBRA", "MÃO-DE-OBRA", "MÃO DE OBRA"),
    ("POÇO-DE-VISITA PV-01", "POÇO-DE-VISita".upper(), "POÇO DE VISITA"),
    ("ENSAIO DE CORPO-DE-PROVA", "CORPO-DE-PROVA", "CORPO DE PROVA"),
    ("ÁREA DE BOTA FORA", "BOTA FORA", "BOTA-FORA"),
    ("ALINHAMENTO DO MEIO FIO", "MEIO FIO", "MEIO-FIO"),
    ("CONCRETO AUTO-ADENSÁVEL", "AUTO-ADENSÁVEL", "AUTOADENSÁVEL"),
    ("OBRAS DE INFRA-ESTRUTURA", "INFRA-ESTRUTURA", "INFRAESTRUTURA"),
    ("SESSÃO TRANSVERSAL DO CANAL", "SESSÃO TRANSVERSAL", "SEÇÃO TRANSVERSAL"),
    ("SESSÃO PLENA DE ESCOAMENTO", "SESSÃO PLENA", "SEÇÃO PLENA"),
    ("CUMPRIMENTO TOTAL DO TUBO", "CUMPRIMENTO TOTAL", "COMPRIMENTO TOTAL"),
    ("CUMPRIMENTO DE ANCORAGEM", "CUMPRIMENTO DE ANCORAGEM", "COMPRIMENTO DE ANCORAGEM"),
    ("AFIM DE GARANTIR A ESTANQUEIDADE", "AFIM DE", "A FIM DE"),
    ("HAJA VISTO O RELATÓRIO", "HAJA VISTO", "HAJA VISTA"),
    ("O TUBO ESTA INSTALADO NA VALA", "ESTA INSTALADO", "ESTÁ INSTALADO"),
    ("A REDE ESTA LOCALIZADA NO PASSEIO", "ESTA LOCALIZADA", "ESTÁ LOCALIZADA"),
    ("O REGISTRO ESTA EM OPERAÇÃO", "ESTA EM", "ESTÁ EM"),
    # crase, regência, plurais compostos e confusões semânticas de engenharia
    ("TRECHO À MONTANTE DO PV-01", "À MONTANTE", "A MONTANTE"),
    ("DESCARGA À JUSANTE DA BARRAGEM", "À JUSANTE", "A JUSANTE"),
    ("ALVENARIA EXECUTADA À PRUMO", "À PRUMO", "A PRUMO"),
    ("DIMENSIONADO CONFORME À NBR 6118", "CONFORME À", "CONFORME A"),
    ("ALTERAÇÃO IMPLICA EM CUSTO ADICIONAL", "IMPLICA EM", "IMPLICA"),
    ("FALHA ACARRETA EM FISSURAÇÃO", "ACARRETA EM", "ACARRETA"),
    ("O SISTEMA CONSISTE DE ESTACAS", "CONSISTE DE", "CONSISTE EM"),
    ("DEVE OBEDECER A NBR 5410", "OBEDECER A NBR", "OBEDECER À NBR"),
    ("PROCEDER A EXECUÇÃO DA BASE", "PROCEDER A EXECUÇÃO", "PROCEDER À EXECUÇÃO"),
    ("IMÓVEL SITUADO À RUA DAS FLORES", "SITUADO À RUA", "SITUADO NA RUA"),
    ("TRATAM-SE DE TUBULAÇÕES DE PEAD", "TRATAM-SE DE", "TRATA-SE DE"),
    ("INSTALAR OS GUARDAS-CORPOS", "GUARDAS-CORPOS", "GUARDA-CORPOS"),
    ("ALTURAS DOS PÉ-DIREITOS", "PÉ-DIREITOS", "PÉS-DIREITOS"),
    ("LOCAÇÃO DAS BOCA-DE-LOBOS", "BOCA-DE-LOBOS", "BOCAS-DE-LOBO"),
    ("TÍTULO: PLANTA ÁREA GERAL", "PLANTA ÁREA", "PLANTA AÉREA"),
])
def test_expressoes_tecnicas_hifen_homonimos_e_contexto(frase, termo, sugestao):
    """R-ORT-015"""
    achados = dict(ort.problemas_frase(frase))
    assert achados.get(termo) == [sugestao]


def test_expressao_marcada_nao_duplica_ocorrencia_nas_palavras_isoladas(verif):
    """R-ORT-015: se 'PRE-MOLDADO' já foi marcado como expressão, 'PRE' não sai duplicado."""
    doc = ezdxf.new("R2018")
    m = doc.modelspace()
    m.add_text("BLOCO PRE-MOLDADO", height=2).set_placement((10, 10))
    papel = doc.layouts.new("F-01")
    papel.reset_main_viewport()
    papel.add_viewport(center=(150, 100), size=(100, 100), view_center_point=(50, 50), view_height=100)
    doc.layouts.delete("Layout1")
    oc = ort.ocorrencias("obra.dwg", doc, verif)
    assert [o.palavra for o in oc] == ["PRE-MOLDADO"]
    assert oc[0].sugestoes == ["PRÉ-MOLDADO"]


@pytest.mark.parametrize("frase,termo,sugestao", [
    ("TUBO DE PEAD , DN 110", "PEAD ,", "PEAD,"),
    ("VER NOTA 1 ; ITEM 2", "1 ;", "1;"),
    ("TUBO,CONEXÃO E REGISTRO", "TUBO,CONEXÃO", "TUBO, CONEXÃO"),
    ("NOTA DE PROJETO.. VER DETALHE", "..", "."),
    ("TUBO,, CONEXÃO", ",,", ","),
    ("PLANTA BAIXA (ESC. 1:50", "(", "fechar ')'"),
    ("DETALHE [VER PRANCHA 02", "[", "fechar ']'"),
    ("PROJETO DE FUNDAÇÕES E\nE ESTRUTURAS DE CONCRETO", "E\nE", "E"),
    ("TUBO DE DE PEAD DN 110", "DE DE", "DE"),
])
def test_pontuacao_mecanica_detecta_erros_sem_afetar_decimais_nem_reticencias(frase, termo, sugestao):
    """R-ORT-016"""
    achados = dict(ort.problemas_frase(frase))
    assert achados.get(termo) == [sugestao]
    # Decimais, escalas e reticências legítimas não disparam erro
    assert ort.problemas_frase("COTA 12,50 m - ESC. 1:50 - CONTINUA...") == []



# ------------------------------------- decimal, expoente e texto de modelo
def test_so_pontuacao_nao_e_conferido():
    """R-ORT-016: campo vazio de carimbo preenchido com '.' ou '..'."""
    assert ort.problemas_frase("..") == [] and ort.problemas_frase(".") == []


@pytest.mark.parametrize("texto,achado", [
    ("D=16.58m", ("16.58", "16,58")),
    ("fck: 35.0MPa", ("35.0", "35,0")),
    ("Fôrmas da Fundação (Nível -0.50)", ("0.50", "0,50")),
])
def test_decimal_com_ponto_vira_virgula(texto, achado):
    """R-ORT-018"""
    assert [(t, s[0]) for t, s in ort.problemas_numero(texto, "TXT02", [])] == [achado]


@pytest.mark.parametrize("texto", [
    "E(X)= 298.400", "ÁREA = 6.579,56", "1 : 1.000", "COTA 12,50 m", "10.10.2026",
    "4 N8 ø12.5 C=540", "Ø (mm)\t5.0   6.3   8.0",
])
def test_milhar_data_e_bitola_nao_sao_decimal_errado(texto):
    """R-ORT-018"""
    assert ort.problemas_numero(texto, "TXT02", []) == []


def test_layer_ignorado_fica_fora_do_decimal():
    """R-ORT-018"""
    padroes = ["*curva*", "$*", "margem"]
    assert ort.problemas_numero("632.5", "_CURVAS003", padroes) == []
    assert ort.problemas_numero("D=16.58m", "$PROJETO-BASE", padroes) == []
    assert ort.problemas_numero("0.1", "MARGEM", padroes) == []
    assert ort.problemas_numero("D=16.58m", "TXT02", padroes)


def test_lista_padrao_de_layers_ignorados():
    """R-ORT-018"""
    padroes = ort.le_padroes(ort.IGNORAR_DECIMAL)
    assert {"*curva*", "*grid*", "$*", "formato", "margem"} <= set(padroes)


@pytest.mark.parametrize("texto,achado", [
    ("FÔRMAS (m2)", ("m2", "m²")),
    ("360kg/m3", ("m3", "m³")),
    ("ÁREA = 12,5 cm2", ("cm2", "cm²")),
])
def test_expoente_de_unidade(texto, achado):
    """R-ORT-019"""
    assert achado in [(t, s[0]) for t, s in ort.problemas_numero(texto, "TXT02", [])]


def test_expoente_respeita_os_layers_ignorados():
    """R-ORT-019: área de lote da base ($PROJETO-BASE) é topografia, fica de fora."""
    assert ort.problemas_numero("537,51m2", "$PROJETO-BASE", ["$*"]) == []
    assert ort.problemas_numero("537,51m2", "TXT02", ["$*"])


def test_vazao_com_tempo_fica_para_a_regra_de_unidade():
    """R-ORT-019: m3/h já é da R-ORT-014; não marca duas vezes."""
    assert [t for t, _ in ort.problemas_frase("5 m3/h")] == ["5 m3/h"]


def test_texto_de_modelo_esquecido():
    """R-ORT-020"""
    achados = [(t, s[0]) for t, s in ort.problemas_frase("3 N1 øX.X c/xx - ??? - A DEFINIR")]
    assert ("øX.X", "preencher") in achados and ("xx", "preencher") in achados
    assert ("???", "preencher") in achados and ("A DEFINIR", "preencher") in achados
    assert ort.problemas_frase("RESERVATÓRIO 8,00X8,00M") == []


def test_ocorrencias_trazem_decimal_respeitando_o_layer():
    """R-ORT-018: ligado ao fluxo (o layer de cada texto decide)."""
    doc = _doc()
    doc.layers.add("_CURVAS003")
    m = doc.modelspace()
    m.add_text("D=16.58m", height=2, dxfattribs={"layer": "TXT"}).set_placement((20, 60))
    m.add_text("632.5", height=2, dxfattribs={"layer": "_CURVAS003"}).set_placement((20, 70))
    v = ort.Verificador(ignorar_decimal=["*curva*"])
    achados = {o.palavra: o.sugestoes for o in ort.ocorrencias("x.dwg", doc, v)}
    assert achados.get("16.58") == ["16,58"] and "632.5" not in achados


@pytest.mark.parametrize("texto", ["NBR 6118/2014  Tab. 9.1", "conforme item 7.4", "Tabela 2.3",
                                   "R (cm)\t2.0   3.5   4.0   5.0"])
def test_referencia_de_norma_e_tabela_de_aco_nao_sao_decimal(texto):
    """R-ORT-018"""
    assert ort.problemas_numero(texto, "TEXTO", []) == []


def test_letras_espacadas_nao_sao_palavra_repetida():
    """R-ORT-016. Regressão: 'R U A   A N T O N I O' (nome de rua espaçado) virava 'A A'."""
    assert not any("A" == s[0] for _, s in ort.problemas_frase("R U A          A N T O N I O"))
    assert ("E \n E", ["E"]) in ort.problemas_frase("FUNDAÇÕES E \n E ESTRUTURAS")
    assert any(s == ["E"] for _, s in ort.problemas_frase("FUNDAÇÕES E E ESTRUTURAS"))


def test_texto_de_modelo_com_bitola_nao_vira_palavra_suspeita():
    """R-ORT-020: o pedaço 'øX' é do modelo, não vai para o corretor."""
    doc = _doc()
    doc.modelspace().add_text("3 N1 øX.X c/xx", height=2, dxfattribs={"layer": "TXT"}).set_placement((20, 60))
    oc = [o.palavra for o in ort.ocorrencias("x.dwg", doc, ort.Verificador(ignorar_decimal=[]))
          if "X" in o.texto]
    assert "øX" not in oc and "øX.X" in oc
